"""Pack de préparation MACF : CSV, XLSX et synthèse PDF.

Chaque sortie porte la mention « préparation — à vérifier par le déclarant MACF autorisé », la phrase de
renvoi et l'avertissement général. Les valeurs lues dans les documents sont des données : une cellule qui
commence par ``=``, ``+``, ``-`` ou ``@`` est neutralisée (injection de formule).
"""

from __future__ import annotations

import csv
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from controldone.formatage import format_nombre
from controldone.guardrails import AVERTISSEMENT, PHRASE_RENVOI, assert_clean
from controldone.macf.agregation import Agregat, SyntheseSeuil, agreger, synthese_seuil
from controldone.macf.codes import ListeCodesMACF, charger_liste
from controldone.macf.demandes import DONNEES_A_DEMANDER, MENTION_PREPARATION
from controldone.macf.selection import LigneMACF

__all__ = [
    "COLONNES_LIGNES",
    "PackMACF",
    "ecrire_csv",
    "ecrire_pack",
    "ecrire_pdf",
    "ecrire_xlsx",
    "preparer_pack",
]

NE_FAIT_PAS = (
    "ControlDOne ne dépose aucune déclaration MACF, ne dit pas si le MACF s'applique à vos importations et ne "
    "se prononce ni sur le classement des marchandises ni sur les valeurs d'émissions. Ce pack rassemble des "
    "données lues sur vos documents pour vous ou votre déclarant MACF autorisé."
)

COLONNES_LIGNES = [
    "Dossier",
    "MRN",
    "Date d'acceptation",
    "Période",
    "Article",
    "Code imprimé",
    "Description imprimée",
    "Correspondance liste",
    "Code de la liste",
    "Secteur",
    "Pays d'origine imprimé",
    "Masse nette lue (kg)",
    "Masse nette imprimée",
    "Fournisseur (tel qu'imprimé)",
    "Installation",
    "Page",
    "Vérification",
]
COLONNES_AGREGATS = [
    "Période",
    "Code imprimé",
    "Pays d'origine",
    "Fournisseur",
    "Installation",
    "Secteur",
    "Lignes",
    "Masse nette lue (kg)",
    "Lignes sans masse",
    "Dossiers",
]
LIBELLES_STATUT = {
    "dans_liste": "code imprimé dans la liste — à faire vérifier",
    "a_preciser": "code imprimé incomplet — à préciser et à faire vérifier",
}


@dataclass
class PackMACF:
    client: str
    annee: int
    lignes: list[LigneMACF]
    agregats: list[Agregat]
    seuil: SyntheseSeuil
    liste: ListeCodesMACF
    date_preparation: date
    mention: str = MENTION_PREPARATION
    phrase_renvoi: str = PHRASE_RENVOI
    avertissement: str = AVERTISSEMENT
    textes: list[str] = field(default_factory=list)


def preparer_pack(
    lignes: Iterable[LigneMACF],
    *,
    annee: int,
    client: str,
    date_preparation: date,
    liste: ListeCodesMACF | None = None,
) -> PackMACF:
    liste = liste or charger_liste()
    tout = list(lignes)
    ls = [li for li in tout if li.annee == annee]
    seuil = synthese_seuil(tout, annee, seuil_t=liste.seuil_tonnes)
    textes = [MENTION_PREPARATION, NE_FAIT_PAS, *seuil.textes]
    for t in textes:
        assert_clean(t)
    return PackMACF(
        client=client,
        annee=annee,
        lignes=ls,
        agregats=agreger(ls, annee=annee),
        seuil=seuil,
        liste=liste,
        date_preparation=date_preparation,
        textes=textes,
    )


def _statut_liste(pack: PackMACF) -> str:
    return "à vérifier" if pack.liste.a_verifier else pack.liste.statut


def _neutre(x: Any) -> Any:
    if isinstance(x, str) and x[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + x
    return x


def _ligne(li: LigneMACF) -> list[Any]:
    return [
        li.dossier_reference or li.dossier_id,
        li.mrn,
        li.date_acceptation.isoformat() if li.date_acceptation else None,
        li.periode,
        li.numero_article,
        li.code_imprime,
        li.description,
        LIBELLES_STATUT.get(li.statut_code.value),
        li.code_liste,
        li.secteur_libelle,
        li.pays_origine,
        str(li.masse_nette_kg) if li.masse_nette_kg is not None else None,
        li.masse_nette_brut,
        li.fournisseur,
        li.installation or "à demander au fournisseur",
        li.page,
        li.verification,
    ]


def _agregat(a: Agregat) -> list[Any]:
    return [
        a.periode,
        a.code_imprime,
        a.pays_origine,
        a.fournisseur,
        a.installation,
        a.secteur,
        a.nombre_lignes,
        str(a.masse_nette_kg),
        a.lignes_sans_masse,
        ", ".join(a.dossiers),
    ]


def _entete(pack: PackMACF) -> list[list[Any]]:
    return [
        [MENTION_PREPARATION.upper()],
        [f"Client : {pack.client}"],
        [f"Année : {pack.annee}"],
        [f"Liste des codes : version {pack.liste.version} ({_statut_liste(pack)}) — {pack.liste.source_url}"],
        [NE_FAIT_PAS],
        [PHRASE_RENVOI],
        [AVERTISSEMENT],
        [],
    ]


def ecrire_csv(pack: PackMACF, dossier: Path | str) -> list[Path]:
    """``macf_lignes_<annee>.csv`` et ``macf_agregats_<annee>.csv`` (UTF-8 avec BOM, séparateur « ; »)."""
    d = Path(dossier)
    d.mkdir(parents=True, exist_ok=True)
    sorties = []
    for nom, colonnes, rangs in (
        ("lignes", COLONNES_LIGNES, [_ligne(li) for li in pack.lignes]),
        ("agregats", COLONNES_AGREGATS, [_agregat(a) for a in pack.agregats]),
    ):
        p = d / f"macf_{nom}_{pack.annee}.csv"
        with open(p, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            for r in _entete(pack):
                w.writerow([_neutre(x) for x in r])
            w.writerow(colonnes)
            for r in rangs:
                w.writerow([_neutre(x) for x in r])
        sorties.append(p)
    return sorties


def ecrire_xlsx(pack: PackMACF, chemin: Path | str) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    from controldone.rapport.export import neutraliser_formules

    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    for r in _entete(pack):
        ws.append(r)
    for t in pack.seuil.textes:
        ws.append([t])
    ws["A1"].font = Font(bold=True, color="B00020")
    ws.append([])
    ws.append(["Données à demander aux fournisseurs"])
    for dd in DONNEES_A_DEMANDER:
        ws.append([dd.libelle, dd.precision])
    for titre, colonnes, rangs in (
        ("Lignes", COLONNES_LIGNES, [_ligne(li) for li in pack.lignes]),
        ("Agrégats", COLONNES_AGREGATS, [_agregat(a) for a in pack.agregats]),
    ):
        f = wb.create_sheet(titre)
        f.append([MENTION_PREPARATION.upper()])
        f.append(colonnes)
        for c in f[2]:
            c.font = Font(bold=True)
        for r in rangs:
            f.append([_neutre(x) for x in r])
    for f in wb.worksheets:
        neutraliser_formules(f)
    p = Path(chemin)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return p


def ecrire_pdf(pack: PackMACF, chemin: Path | str) -> Path:
    """Synthèse PDF (A4) : mention en bandeau sur chaque page, cumul, agrégats, liste à demander."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    base = ParagraphStyle("b", fontName="Helvetica", fontSize=8.5, leading=11.5)
    h1 = ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=15, leading=19, spaceAfter=6)
    h2 = ParagraphStyle(
        "h2", parent=base, fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=4
    )
    petit = ParagraphStyle("p", parent=base, fontSize=7, leading=9)

    def p(t: Any, s: ParagraphStyle = base) -> Paragraph:
        return Paragraph(escape("" if t is None else str(t)).replace("\n", "<br/>"), s)

    def bandeau(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#B00020"))
        canvas.setFont("Helvetica-Bold", 9)
        canvas.drawCentredString(A4[0] / 2, A4[1] - 12 * mm, MENTION_PREPARATION.upper())
        canvas.setFillColor(colors.grey)
        canvas.setFont("Helvetica", 6.5)
        canvas.drawString(15 * mm, 10 * mm, f"ControlDOne — pack MACF {pack.annee} — page {doc.page}")
        canvas.restoreState()

    elems: list[Any] = [
        p(f"Préparation des données MACF (CBAM) — année {pack.annee}", h1),
        p(f"Client : {pack.client} — préparé le {pack.date_preparation.isoformat()}"),
        Spacer(1, 4),
        p(NE_FAIT_PAS),
        Spacer(1, 4),
        p("Cumul annuel et seuil de 50 t (arithmétique)", h2),
        *[p(t) for t in pack.seuil.textes],
        p("Masses nettes par code, origine, fournisseur et période", h2),
    ]
    entete = ["Période", "Code imprimé", "Origine", "Fournisseur", "Lignes", "Masse (kg)"]
    donnees = [entete] + [
        [
            a.periode,
            a.code_imprime,
            a.pays_origine,
            p(a.fournisseur, petit),
            a.nombre_lignes,
            format_nombre(a.masse_nette_kg, 3),
        ]
        for a in pack.agregats
    ]
    if len(donnees) == 1:
        donnees.append(["—", "aucune ligne retenue", "", "", "", ""])
    t = Table(donnees, colWidths=[18 * mm, 30 * mm, 16 * mm, 66 * mm, 14 * mm, 26 * mm], repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
                ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF4")),
                ("ALIGN", (4, 1), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    elems += [t, p("Données à demander aux fournisseurs", h2)]
    elems += [p(f"• {dd.libelle} : {dd.precision}") for dd in DONNEES_A_DEMANDER]
    elems += [
        p("Sources et hypothèses", h2),
        p(
            f"Liste des codes NC : {pack.liste.source_texte} — version {pack.liste.version}, statut "
            f"« {_statut_liste(pack)} », consultée le {pack.liste.consulte_le} ({pack.liste.source_url})."
        ),
        p(f"Seuil de 50 t : {pack.liste.seuil_source_url}."),
        p(
            "Correspondance faite sur le code imprimé de la déclaration, sans avis de classement ; chaque ligne est "
            "à faire vérifier."
        ),
        Spacer(1, 6),
        p(PHRASE_RENVOI),
        Spacer(1, 4),
        p(AVERTISSEMENT, petit),
    ]
    out = Path(chemin)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(out),
        pagesize=A4,
        topMargin=20 * mm,
        bottomMargin=16 * mm,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        title=f"Pack MACF {pack.annee} — {MENTION_PREPARATION}",
        author="ControlDOne",
    )
    doc.build(elems, onFirstPage=bandeau, onLaterPages=bandeau)
    return out


def ecrire_pack(pack: PackMACF, dossier: Path | str) -> dict[str, Path | list[Path]]:
    """Écrit CSV, XLSX et PDF dans ``dossier``."""
    d = Path(dossier)
    return {
        "csv": ecrire_csv(pack, d),
        "xlsx": ecrire_xlsx(pack, d / f"macf_pack_{pack.annee}.xlsx"),
        "pdf": ecrire_pdf(pack, d / f"macf_synthese_{pack.annee}.pdf"),
    }

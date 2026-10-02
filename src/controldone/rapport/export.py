"""Exports du rapport : JSON (``controldone.rapport/1.0.0``) et tableur XLSX (SPEC §18.1).

Tableur : onglets synthèse, dossiers, constats, contrôles, documents, non lus, méthode, avec des colonnes
vides de revue humaine (décision, commentaire, relu par, date).
"""

from __future__ import annotations

import json
from collections.abc import Sequence
from decimal import Decimal
from pathlib import Path
from typing import Any

from controldone.findings_io import findings_json
from controldone.guardrails import AVERTISSEMENT
from controldone.pipeline import ResultatDossier
from controldone.rapport.vue import VERSION_RAPPORT, RapportVue

__all__ = ["SCHEMA_RAPPORT", "ecrire_xlsx", "findings_lot_json", "neutraliser_formules", "rapport_json"]

SCHEMA_RAPPORT = f"controldone.rapport/{VERSION_RAPPORT}"
COLONNES_REVUE = ["Décision (valider / rejeter / à vérifier)", "Commentaire", "Relu par", "Date de revue"]


def _d(x: Decimal | None) -> str | None:
    return None if x is None else str(x)


def rapport_json(vue: RapportVue, resultats: Sequence[ResultatDossier]) -> dict[str, Any]:
    """Source de vérité de l'interface : synthèse + ``findings`` de chaque dossier (Annexe C)."""
    return {
        "schema": SCHEMA_RAPPORT,
        "client": vue.client,
        "offre": vue.offre,
        "periode": vue.periode,
        "date": vue.date,
        "donnees_fictives": vue.demo,
        "versions": vue.versions,
        "executions": vue.execution_ids,
        "empreinte_tolerances": vue.empreinte,
        "dossiers_figes": [{"dossier_id": rd.dossier.id, "reference": rd.dossier.reference,
                            "version": rd.dossier.version} for rd in resultats],
        "synthese": {
            "statuts": {code: n for _lib, code, n in vue.statuts},
            "recouvrable_certain": vue.recouvrable_certain,
            "recouvrable_a_verifier": vue.recouvrable_a_verifier,
            "ecarts_documentaires": {"nombre": vue.ecarts_documentaires_nb, "montant_absolu": vue.ecarts_documentaires},
            "ecarts_calcul_declaration": {"nombre": vue.ecarts_calcul_nb, "montant_absolu": vue.ecarts_calcul},
            "points_professionnel": vue.nb_renvois,
            "mention": vue.mention_validation,
        },
        "prochaines_actions": [{"priorite": a.priorite, "titre": a.titre, "details": a.details} for a in vue.actions],
        "non_lus": [{"fichier": f, "motif": m, "pages": p} for f, m, p in vue.non_lus],
        "dossiers": [json.loads(findings_json(rd.findings)) for rd in resultats],
        "avertissement": AVERTISSEMENT,
    }


def findings_lot_json(resultats: Sequence[ResultatDossier]) -> list[dict[str, Any]]:
    return [json.loads(findings_json(rd.findings)) for rd in resultats]


def neutraliser_formules(ws: Any) -> int:
    """Toute cellule texte qu'openpyxl a typée « formule » (chaîne commençant par « = ») redevient du texte.

    Les valeurs viennent des documents déposés (numéros, références, noms de fichiers) : une chaîne
    ``=HYPERLINK(…)`` ou ``=cmd|…`` ne doit jamais s'exécuter dans le tableur de celui qui ouvre l'export
    (injection de formule, revue de sécurité RS-11). La valeur affichée reste identique."""
    n = 0
    for ligne in ws.iter_rows():
        for cellule in ligne:
            if cellule.data_type == "f" and isinstance(cellule.value, str):
                cellule.data_type = "s"
                n += 1
    return n


def ecrire_xlsx(vue: RapportVue, resultats: Sequence[ResultatDossier], chemin: Path | str) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    entete_font = Font(bold=True, color="FFFFFF")
    entete_fill = PatternFill("solid", fgColor="1D3557")
    revue_fill = PatternFill("solid", fgColor="FFF4E0")

    def feuille(titre: str, entetes: list[str], lignes: list[list[Any]], *, revue: bool = False, premiere=False):
        ws = wb.active if premiere else wb.create_sheet()
        ws.title = titre
        cols = entetes + (COLONNES_REVUE if revue else [])
        ws.append(cols)
        for i, _c in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=i)
            cell.font, cell.fill = entete_font, entete_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if revue and i > len(entetes):
                cell.fill = PatternFill("solid", fgColor="9A5B00")
        for li in lignes:
            ws.append(list(li) + ([""] * len(COLONNES_REVUE) if revue else []))
        for i, c in enumerate(cols, start=1):
            largeur = max([len(str(c))] + [len(str(x[i - 1])) for x in lignes if i - 1 < len(x)] + [8])
            ws.column_dimensions[get_column_letter(i)].width = min(60, largeur + 2)
            if revue and i > len(entetes):
                for r in range(2, len(lignes) + 2):
                    ws.cell(row=r, column=i).fill = revue_fill
        ws.freeze_panes = "A2"
        neutraliser_formules(ws)
        return ws

    synth = [["Client", vue.client], ["Période", vue.periode], ["Offre", vue.offre], ["Date", vue.date],
             ["Données fictives", "oui" if vue.demo else "non"], ["Dossiers", vue.nb_dossiers]]
    synth += [[f"Dossiers — {lib}", n] for lib, _c, n in vue.statuts]
    synth += [["Montant recouvrable certain", vue.recouvrable_certain],
              ["Montant recouvrable à vérifier (jamais additionné au précédent)", vue.recouvrable_a_verifier],
              ["Écarts de valeur entre documents (nombre)", vue.ecarts_documentaires_nb],
              ["Écarts de valeur entre documents (montant absolu)", vue.ecarts_documentaires],
              ["Écarts de calcul sur la déclaration (nombre)", vue.ecarts_calcul_nb],
              ["Écarts de calcul sur la déclaration (montant absolu)", vue.ecarts_calcul],
              ["Points à faire vérifier par un professionnel", vue.nb_renvois],
              ["Mention", vue.mention_validation], ["Avertissement", vue.avertissement]]
    feuille("Synthèse", ["Indicateur", "Valeur"], synth, premiere=True)
    feuille("Dossiers", ["Dossier", "Facture transitaire", "Transport", "MRN", "Facture commerciale", "TVA acheteur",
                         "TVA importateur", "Montant facturé", "Montant déclaré", "Statut", "Raisons",
                         "Recouvrable certain (EUR)", "Recouvrable à vérifier (EUR)"],
            [[d.reference, *[v for _k, v in d.cles], d.tva_acheteur, d.tva_importateur, d.montant_facture,
              d.montant_declare, d.statut, d.raisons, float(d.recouvrable_certain), float(d.recouvrable_a_verifier)]
             for d in vue.dossiers], revue=True)
    constats = []
    for rd in resultats:
        for r in rd.resultats:
            c = r.constat
            if c is None:
                continue
            constats.append([rd.dossier.reference, c.id, c.controle_id, c.niveau.value, c.nature_montant.value,
                             c.composante.value if c.composante else "", _d(c.montant_en_jeu) or "",
                             "oui" if c.renvoi else "non", ", ".join(x.value for x in c.raisons),
                             c.libelle if c.motif_blocage is None else "(libellé retenu pour relecture)",
                             c.prochaine_action if c.motif_blocage is None else "",
                             _d(r.tolerance_appliquee) or "", _d(r.seuil_certitude_applique) or "",
                             " | ".join(f"{p.role.value}: {p.valeur_brute or p.calcul or ''} (p. {p.page or '-'})"
                                        for p in c.preuves)])
    feuille("Constats", ["Dossier", "Constat", "Contrôle", "Niveau", "Nature du montant", "Composante",
                         "Montant en jeu (EUR)", "Renvoi", "Raisons", "Libellé", "Prochaine action", "Tolérance",
                         "Seuil de certitude", "Preuves"], constats, revue=True)
    feuille("Contrôles", ["Dossier", "Contrôle", "Sous-contrôle", "Unité", "Résultat", "Raison", "Attendu",
                          "Constaté", "Écart", "Tolérance"],
            [[rd.dossier.reference, r.controle_id, r.sous_controle or "", r.unite, r.outcome.value,
              r.raison_code.value if r.raison_code else "", r.attendu or "", r.constate or "", _d(r.ecart) or "",
              _d(r.tolerance_appliquee) or ""] for rd in resultats for r in rd.resultats])
    docs = []
    for rd in resultats:
        for lien in rd.dossier.liens:
            d = rd.documents.get(lien.document_id)
            if d is None:
                continue
            fic = rd.fichiers.get(d.pages[0].fichier_id) if d.pages else None
            docs.append([rd.dossier.reference, d.id, d.type.value, d.sous_type or "",
                         fic.chemin_relatif if fic else "", ", ".join(str(p.numero) for p in d.pages),
                         d.confiance_classement, lien.role.value, lien.force.value,
                         ", ".join(s.value for s in lien.signaux), lien.score])
    feuille("Documents", ["Dossier", "Document", "Type", "Sous-type", "Fichier", "Pages", "Confiance classement",
                          "Rôle", "Rattachement", "Signaux", "Score"], docs, revue=True)
    feuille("Non lus", ["Fichier", "Motif", "Pages"], [list(x) for x in vue.non_lus], revue=True)
    meth = [[x, ""] for x in vue.limites] + [[k, v] for k, v in vue.tolerances]
    meth += [["Empreinte des tolérances", vue.empreinte]] + [[f"Version {k}", v] for k, v in vue.versions.items()]
    meth += [[f"Extracteur {k}", v] for k, v in vue.extracteurs] + [["Modèle de langage", vue.modele_llm],
                                                                     ["Avertissement", vue.avertissement]]
    feuille("Méthode", ["Élément", "Valeur"], meth)
    p = Path(chemin)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return p

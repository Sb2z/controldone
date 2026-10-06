"""Fichier -> pages avec texte (SPEC §5.1, §6.2.5, §7.3, §20.3).

- PDF : texte natif (pdfplumber, mots positionnés). Le texte **invisible** (mode de rendu invisible,
  blanc sans fond sombre, micro-police) est retiré du texte de la page et gardé à part pour l'audit
  (§20.2). Qualité : ``natif`` si ≥ 85 % des caractères forment des mots plausibles ou des nombres ;
  sinon OCR et l'on garde la meilleure des deux sources (``natif_faible`` si le natif l'emporte).
- OCR : rendu pypdfium2 à 300 dpi, prétraitement (``ingest.pretraitement`` : gris par maximum des canaux, traits
  de télécopie effacés, contraste étiré, médian sur télécopie), orientation (OSD + essai des autres orientations
  si la lecture est mauvaise), feuille « deux pages par feuille » lue moitié par moitié, désinclinaison,
  Tesseract ``fra+eng``, mots et confiances conservés, ``score_ocr`` (D-2601 à D-2606).
- Images PNG/JPEG/TIFF multipage : OCR de chaque image.
- Tableurs XLSX/ODS et CSV : une page par feuille, cellules jointes par `` | `` ligne par ligne.
- XML : une page, texte brut conservé. Corps de courriel : une page de texte.

Le travail sur les octets du fichier (analyse PDF, rendu, OCR) tourne dans un **processus séparé** par fichier,
avec limite de temps et de mémoire (issu d'un forkserver préchargé ; à défaut ``_worker``). Rien n'est persisté ici, sauf le cache disque facultatif
(clé : ``sha256 + n° page + version``).
"""

from __future__ import annotations

import hashlib
import io
import itertools
import json
import logging
import multiprocessing
import os
import re
import subprocess
import sys
import tempfile
import threading
import zipfile
from collections.abc import Sequence
from dataclasses import asdict, dataclass, replace
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from itertools import pairwise
from pathlib import Path

from controldone.ids import IdGenerator, Prefixe, nouvel_id
from controldone.model import Fichier, Page
from controldone.model.enums import QualiteTexte

from .sniff import MIME_CSV, MIME_EML, MIME_ODS, MIME_PDF, MIME_TEXTE, MIME_XLSX, MIME_XML, decoder_texte
from .texte import Ligne, Mot, PageText, construire_lignes, score_texte

log = logging.getLogger("controldone.ingest.pages")

__all__ = [
    "CIBLES_ISOLEES",
    "MAX_PIXELS_RENDU_OCR",
    "VERSION_PAGES",
    "CachePagesDisque",
    "OptionsPages",
    "PageExtraite",
    "echelle_rendu_ocr",
    "executer_isole",
    "extraire_pages",
    "extraire_pages_local",
    "ocr_disponible",
    "textes_pages",
    "version_ocr",
]

#: Version de l'algorithme de pages (entre dans la clé d'idempotence §7 étape 2 avec Tesseract).
VERSION_PAGES = (
    "1.1.1"  # 1.1.0 : prétraitement OCR, deux pages par feuille, réessai d'orientation (D-2601 à D-2606)
)
# 1.1.1 : feuille « deux pages » écartée quand une ligne de texte touche la coupure (D-2902)

SEUIL_NATIF = 0.85
SEUIL_NATIF_FAIBLE = 0.5
MIN_CARACTERES_NATIF = 25
SEUIL_ILLISIBLE_OCR = 0.45
#: Bornes du texte **positionné** d'une page de texte, de tableur ou de CSV (D-1604) : au-delà, le texte de la
#: page reste complet mais les mots positionnés s'arrêtent (avertissement ``texte_positionne_tronque``). Un XML ou
#: un CSV de 50 Mo produisait sinon des millions de mots (plusieurs Go dans le processus de pages, puis dans le
#: processus principal qui les relit). Les documents réels en sont très loin (corpus : moins de 300 lignes).
MAX_LIGNES_POSITIONNEES = 50_000
MAX_MOTS_POSITIONNES = 500_000
#: Bornes de lecture d'une feuille de tableur (colonnes lues, cellules lues par feuille).
MAX_COLONNES_TABLEUR = 512
MAX_CELLULES_FEUILLE = 2_000_000


@dataclass(frozen=True)
class OptionsPages:
    ocr: bool = True
    dpi: int = 300
    langues: str = "fra+eng"
    #: Forcer l'OCR même si le texte natif est bon (diagnostic).
    forcer_ocr: bool = False
    #: Temps alloué au processus isolé : ``base + par_page × pages``.
    timeout_base_s: float = 30.0
    timeout_par_page_s: float = 60.0
    memoire_mo: int = 3072
    #: Exécution dans un processus séparé (recommandé ; ``False`` pour déboguer).
    isoler: bool = True
    cache_dir: str | None = None
    max_pages: int = 300

    def empreinte(self) -> str:
        cle = f"{self.dpi}|{self.langues}|{self.ocr}|{self.forcer_ocr}"
        return hashlib.sha256(cle.encode()).hexdigest()[:8]


@dataclass
class PageExtraite:
    """Page du modèle (persistée par l'assemblage) + texte positionné (pour les extracteurs)."""

    page: Page
    texte: PageText


@lru_cache(maxsize=1)
def version_ocr() -> str:
    """Version de Tesseract (``indisponible`` si absent)."""
    try:
        import pytesseract

        return str(pytesseract.get_tesseract_version()).split()[0]
    except Exception:
        return "indisponible"


def ocr_disponible() -> bool:
    return version_ocr() != "indisponible"


def _cle_version(options: OptionsPages) -> str:
    return f"{VERSION_PAGES}+tess{version_ocr()}+{options.empreinte()}"


# --- cache disque -------------------------------------------------------------------------------------


class CachePagesDisque:
    """Cache ``PageText`` sur disque : un fichier JSON par (sha256, page, version).

    Le texte y est **en clair** (fichiers 0600, répertoires 0700) : réservé au banc et au développement ; le
    découpeur du service l'ignore en production (``CONTROLDONE_ENV=prod``, revue de sécurité RS-03)."""

    def __init__(self, dossier: str | Path) -> None:
        self.dossier = Path(dossier)

    def _chemin(self, sha: str, version: str, nom: str) -> Path:
        v = hashlib.sha256(version.encode()).hexdigest()[:12]
        return self.dossier / sha[:2] / f"{sha}_{v}_{nom}.json"

    def lire(self, sha: str, version: str) -> list[PageText] | None:
        idx = self._chemin(sha, version, "index")
        try:
            n = json.loads(idx.read_text("utf-8"))["pages"]
            return [
                PageText.from_dict(json.loads(self._chemin(sha, version, f"p{i}").read_text("utf-8")))
                for i in range(1, n + 1)
            ]
        except (OSError, ValueError, KeyError, TypeError):
            return None

    def ecrire(self, sha: str, version: str, pages: list[PageText]) -> None:
        try:
            for p in pages:
                c = self._chemin(sha, version, f"p{p.numero}")
                c.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
                _ecrire_atomique(c, json.dumps(p.to_dict(), ensure_ascii=False))
            _ecrire_atomique(self._chemin(sha, version, "index"), json.dumps({"pages": len(pages)}))
        except OSError:
            pass  # le cache est une optimisation


def _ecrire_atomique(chemin: Path, contenu: str) -> None:
    """Écriture atomique, fichier privé (0600) : le cache contient du texte de document en clair."""
    tmp = chemin.with_suffix(f".{os.getpid()}.tmp")
    fd = os.open(tmp, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        f.write(contenu)
    os.chmod(tmp, 0o600)
    tmp.replace(chemin)


# --- API publique -------------------------------------------------------------------------------------


def extraire_pages(
    contenu: bytes,
    *,
    fichier: Fichier | None = None,
    type_mime: str | None = None,
    options: OptionsPages | None = None,
    ids: IdGenerator | None = None,
    textes: list[PageText] | None = None,
) -> list[PageExtraite]:
    """``Fichier`` -> ``Page`` avec texte. Fonction pure (hors cache disque facultatif).

    Ne lève pas pour un fichier difficile : une page illisible porte ``qualite_texte = illisible``.
    ``textes`` : résultat de ``textes_pages`` déjà calculé pour ce fichier (préchargement parallèle d'un lot) ;
    seuls les identifiants de page sont alors attribués ici.
    """
    if textes is None:
        textes = textes_pages(contenu, fichier=fichier, type_mime=type_mime, options=options)
    fid = fichier.id if fichier else "fic_inconnu"
    sortie = []
    for t in textes:
        pid = ids.nouveau(Prefixe.page) if ids is not None else nouvel_id(Prefixe.page)
        page = Page(
            id=pid,
            client_id=fichier.client_id if fichier else None,
            fichier_id=fid,
            numero=t.numero,
            rotation_appliquee=t.rotation,
            qualite_texte=t.qualite,
            score_ocr=None if t.score_ocr is None else round(min(1.0, max(0.0, t.score_ocr)), 4),
            texte=t.texte,
            sha256_texte=hashlib.sha256(t.texte.encode("utf-8")).hexdigest(),
            feuille=t.feuille,
            largeur=t.largeur,
            hauteur=t.hauteur,
        )
        sortie.append(PageExtraite(page=page, texte=t))
    return sortie


def textes_pages(
    contenu: bytes,
    *,
    fichier: Fichier | None = None,
    type_mime: str | None = None,
    options: OptionsPages | None = None,
) -> list[PageText]:
    """Texte des pages d'un fichier (cache disque facultatif, puis processus isolé ou extraction locale).

    Sans identifiant ni état partagé : peut tourner dans plusieurs fils à la fois (un processus isolé chacun)."""
    opts = options or OptionsPages()
    mime = type_mime or (fichier.type_mime if fichier else None)
    if mime is None:
        from .sniff import detecter_type

        mime = detecter_type(contenu)
    sha = fichier.sha256 if fichier else hashlib.sha256(contenu).hexdigest()
    version = _cle_version(opts)
    cache = CachePagesDisque(opts.cache_dir) if opts.cache_dir else None
    textes = cache.lire(sha, version) if cache else None
    if textes is None:
        if opts.isoler:
            n_pages = fichier.nombre_pages if fichier and fichier.nombre_pages else None
            textes = _extraire_isole(contenu, mime, opts, n_pages)
        else:
            textes = extraire_pages_local(contenu, mime, opts)
        if cache and textes and not any("timeout" in a for p in textes for a in p.avertissements):
            cache.ecrire(sha, version, textes)
    return textes


def _extraire_isole(
    contenu: bytes, mime: str, opts: OptionsPages, n_pages: int | None = None
) -> list[PageText]:
    """Exécute ``extraire_pages_local`` dans un processus séparé (temps et mémoire bornés).

    Un processus **neuf par fichier**, issu d'un forkserver préchargé (D-1402) : mêmes limites (RLIMIT_AS, délai,
    arrêt forcé), environnement sans secrets (RS-14), sans le coût d'un nouvel interpréteur (~0,4 s par fichier).
    Repli sur ``python -m controldone.ingest._worker`` si la plateforme n'offre pas de forkserver."""
    n_estime = n_pages or _estimer_pages(contenu, mime)
    timeout = opts.timeout_base_s + opts.timeout_par_page_s * max(1, n_estime)
    options = {**asdict(opts), "isoler": False, "cache_dir": None}
    if _FORKSERVER_DISPONIBLE:
        donnees, motif = _executer_forkserver(contenu, mime, options, opts.memoire_mo, timeout)
    else:
        donnees, motif = _executer_sous_processus(contenu, mime, options, opts.memoire_mo, timeout)
    if donnees is not None:
        return [PageText.from_dict(d) for d in json.loads(donnees)["pages"]]
    # Échec ou délai dépassé : nouvel essai sans OCR (texte natif seulement), puis pages illisibles.
    if opts.ocr:
        pages = _extraire_isole(contenu, mime, replace(opts, ocr=False, timeout_par_page_s=5.0), n_pages)
        for p in pages:
            p.avertissements.append(f"ocr_abandonne:{motif}")
        return pages
    return [_page_illisible(i, motif) for i in range(1, max(1, n_estime) + 1)]


def _executer_sous_processus(
    contenu: bytes, mime: str, options: dict, memoire_mo: int, timeout: float
) -> tuple[str | None, str | None]:
    """Processus isolé par un nouvel interpréteur (``python -m controldone.ingest._worker``)."""
    with tempfile.TemporaryDirectory(prefix="cdo_pages_") as tmp:
        entree = Path(tmp) / "entree.bin"
        entree.write_bytes(contenu)
        sortie = Path(tmp) / "sortie.json"
        req = {"entree": str(entree), "sortie": str(sortie), "mime": mime, "options": options}
        try:
            proc = subprocess.run(
                [sys.executable, "-m", "controldone.ingest._worker"],
                input=json.dumps(req).encode(),
                capture_output=True,
                timeout=timeout,
                env={**_environnement_sans_secrets(), **_ENV_PROCESSUS_PAGES},
                preexec_fn=_limiteur_memoire(memoire_mo) if os.name == "posix" else None,
                check=False,
            )
            if proc.returncode == 0 and sortie.exists():
                return sortie.read_text("utf-8"), None
            return None, f"processus_pages_code_{proc.returncode}"
        except subprocess.TimeoutExpired:
            return None, "timeout"


# --- processus isolé issu d'un forkserver (D-1402) ------------------------------------------------------------

#: Variables fixées dans le processus de pages : un fil par OCR (Tesseract multi-fils se dégrade sous charge).
_ENV_PROCESSUS_PAGES = {"OMP_THREAD_LIMIT": "1", "OPENBLAS_NUM_THREADS": "1"}
#: Modules chargés une fois dans le forkserver (code et bibliothèques seulement : aucune donnée de document).
_PRECHARGES = [
    "controldone.ingest.pages",
    "controldone.ingest.sniff",
    "pdfplumber",
    "pypdfium2",
    "pytesseract",
    "PIL.Image",
    "PIL.ImageStat",
    "openpyxl",
    "lxml.etree",
    "controldone.services.vignettes",
]
_FORKSERVER_DISPONIBLE = os.name == "posix" and "forkserver" in multiprocessing.get_all_start_methods()
_VERROU_FS = threading.Lock()
_PID_FS: list[int | None] = [
    None
]  # processus qui a configuré le forkserver (un fork en hérite sans pouvoir l'utiliser)


class _EnvironnementSansSecrets:
    """Retire les secrets de l'environnement **C** du processus le temps de lancer le forkserver (qui hérite de
    l'environnement C, RS-14). ``os.environ`` (Python) n'est pas modifié : les autres fils n'en voient rien."""

    def __enter__(self):
        propre = _environnement_sans_secrets()  # même filtre que le sous-processus (RS-14)
        self._retires = {k: v for k, v in os.environ.items() if k not in propre}
        self._ajoutes = {k: os.environ.get(k) for k in _ENV_PROCESSUS_PAGES}
        for k in self._retires:
            os.unsetenv(k)
        for k, v in _ENV_PROCESSUS_PAGES.items():
            os.putenv(k, v)

    def __exit__(self, *exc):
        for k, v in self._retires.items():
            os.putenv(k, v)
        for k, v in self._ajoutes.items():
            if v is None:
                os.unsetenv(k)
            else:
                os.putenv(k, v)


def _contexte_forkserver():
    """Contexte ``forkserver`` dont le serveur tourne, lancé avec l'environnement sans secrets. À appeler sous
    ``_VERROU_FS`` avant chaque ``start()`` : un serveur arrêté est relancé ici, jamais par ``start()`` avec
    l'environnement complet."""
    from multiprocessing import forkserver

    serveur = forkserver._forkserver
    if _PID_FS[0] != os.getpid():
        if _PID_FS[0] is not None or serveur._forkserver_pid is not None:
            # Processus issu d'un fork : le serveur hérité est celui du parent (``waitpid`` impossible ici).
            serveur._forkserver_address = serveur._forkserver_alive_fd = serveur._forkserver_pid = None
        _PID_FS[0] = os.getpid()
    ctx = multiprocessing.get_context("forkserver")
    ctx.set_forkserver_preload(_PRECHARGES)
    with _EnvironnementSansSecrets():
        serveur.ensure_running()
    return ctx


if _FORKSERVER_DISPONIBLE:
    from multiprocessing import context as _mp_context
    from multiprocessing import forkserver as _mp_forkserver
    from multiprocessing import popen_forkserver as _mp_popen_forkserver
    from multiprocessing import reduction as _mp_reduction
    from multiprocessing import spawn as _mp_spawn
    from multiprocessing import util as _mp_util

    class _PopenSansMain(_mp_popen_forkserver.Popen):
        """Lancement par le forkserver **sans** réimporter le module principal (``__main__``) du parent.

        Par défaut, chaque enfant d'un forkserver ré-exécute le script ou le module principal du parent (banc,
        serveur web, worker) : coût d'import et effets de bord à chaque fichier. Le processus de pages n'en a pas
        besoin (sa cible est dans ce module) : les clés ``init_main_from_*`` sont retirées des données de
        préparation. Reprend ``popen_forkserver.Popen._launch`` (CPython 3.11 à 3.13)."""

        def _launch(self, process_obj):
            prep = {
                k: v
                for k, v in _mp_spawn.get_preparation_data(process_obj._name).items()
                if k not in ("init_main_from_path", "init_main_from_name")
            }
            buf = io.BytesIO()
            _mp_context.set_spawning_popen(self)
            try:
                _mp_reduction.dump(prep, buf)
                _mp_reduction.dump(process_obj, buf)
            finally:
                _mp_context.set_spawning_popen(None)
            self.sentinel, w = _mp_forkserver.connect_to_new_process(self._fds)
            _parent_w = os.dup(w)
            self.finalizer = _mp_util.Finalize(self, _mp_util.close_fds, (_parent_w, self.sentinel))
            with open(w, "wb", closefd=True) as f:
                f.write(buf.getbuffer())
            self.pid = _mp_forkserver.read_signed(self.sentinel)

    class _ProcessusPages(_mp_context.ForkServerProcess):
        @staticmethod
        def _Popen(process_obj):
            return _PopenSansMain(process_obj)


def _processus_pages(
    envoi, contenu: bytes, mime: str, options: dict, memoire_mo: int, delai_s: float
) -> None:
    """Corps du processus isolé (enfant du forkserver)."""
    import signal

    try:
        nul = os.open(os.devnull, os.O_RDWR)  # rien du document sur les sorties (comme ``capture_output``)
        os.dup2(nul, 1)
        os.dup2(nul, 2)
        os.environ.update(_ENV_PROCESSUS_PAGES)
        _limiteur_memoire(memoire_mo)()
        signal.alarm(int(delai_s) + 30)  # garde-fou si le parent disparaît sans tuer ce processus
        champs = set(OptionsPages.__dataclass_fields__)
        opts = OptionsPages(**{k: v for k, v in options.items() if k in champs})
        pages = extraire_pages_local(contenu, mime, opts)
        envoi.send_bytes(
            json.dumps({"pages": [p.to_dict() for p in pages]}, ensure_ascii=False).encode("utf-8")
        )
        envoi.close()
    except BaseException:
        os._exit(1)


def _executer_forkserver(
    contenu: bytes, mime: str, options: dict, memoire_mo: int, timeout: float
) -> tuple[str | None, str | None]:
    brut, motif = _lancer_forkserver(_processus_pages, (contenu, mime, options, memoire_mo, timeout), timeout)
    return (brut.decode("utf-8") if brut is not None else None), motif


#: Rendus confiés au processus isolé (RS-20, D-3607) : nom -> fonction ``module:attribut`` qui renvoie des octets
#: ou ``None``. Liste fermée : le parent ne peut demander rien d'autre.
CIBLES_ISOLEES = {"rendu_page": "controldone.services.vignettes:rendu_local"}


def _processus_cible(envoi, cible: str, kwargs: dict, memoire_mo: int, delai_s: float) -> None:
    """Corps du processus isolé pour un rendu de ``CIBLES_ISOLEES`` (enfant du forkserver)."""
    import importlib
    import signal

    try:
        nul = os.open(os.devnull, os.O_RDWR)
        os.dup2(nul, 1)
        os.dup2(nul, 2)
        os.environ.update(_ENV_PROCESSUS_PAGES)
        _limiteur_memoire(memoire_mo)()
        signal.alarm(int(delai_s) + 30)
        module, _, attribut = CIBLES_ISOLEES[cible].partition(":")
        resultat = getattr(importlib.import_module(module), attribut)(**kwargs)
        envoi.send_bytes(b"\x00" if resultat is None else b"\x01" + bytes(resultat))
        envoi.close()
    except BaseException:
        os._exit(1)


def executer_isole(
    cible: str, kwargs: dict, *, memoire_mo: int = 1536, delai_s: float = 30.0
) -> tuple[bytes | None, str | None]:
    """Exécute la fonction ``CIBLES_ISOLEES[cible](**kwargs)`` dans un processus isolé (même forkserver que
    l'extraction des pages : environnement sans secrets, ``RLIMIT_AS``, délai, arrêt forcé). Renvoie
    ``(octets, None)``, ``(None, None)`` si la fonction a renvoyé ``None``, ou ``(None, motif)`` en cas d'échec
    (délai, mémoire, plantage de la bibliothèque de rendu). Sans forkserver (plateforme non POSIX) : exécution
    locale, journalisée."""
    import importlib

    if cible not in CIBLES_ISOLEES:
        raise ValueError(f"cible non autorisée : {cible}")
    if not _FORKSERVER_DISPONIBLE:  # pragma: no cover - Linux en production
        log.warning("rendu_non_isole cible=%s", cible)
        module, _, attribut = CIBLES_ISOLEES[cible].partition(":")
        return getattr(importlib.import_module(module), attribut)(**kwargs), None
    brut, motif = _lancer_forkserver(_processus_cible, (cible, kwargs, memoire_mo, delai_s), delai_s)
    if brut is None:
        return None, motif
    return (bytes(brut[1:]) if brut[:1] == b"\x01" else None), None


def _lancer_forkserver(cible_fn, args: tuple, timeout: float) -> tuple[bytes | None, str | None]:
    with _VERROU_FS:
        ctx = _contexte_forkserver()
        recu, envoi = ctx.Pipe(duplex=False)
        proc = _ProcessusPages(target=cible_fn, name="cdo-pages", daemon=True, args=(envoi, *args))
        # Le worker de jobs exécute le pipeline dans un processus « daemon » (jobs.handlers) : multiprocessing y
        # refuse tout enfant parce qu'il ne les attendrait pas à la sortie. Ici l'enfant est toujours attendu ou
        # tué avant de rendre la main (et se termine seul au-delà du délai) : on lève l'interdiction pour ce start.
        config = multiprocessing.current_process()._config
        daemon = config.pop("daemon", None)
        try:
            proc.start()
        finally:
            if daemon is not None:
                config["daemon"] = daemon
    envoi.close()
    donnees = motif = None
    try:
        if recu.poll(timeout):
            try:
                donnees = recu.recv_bytes()
            except (EOFError, OSError):
                donnees = None
        else:
            motif = "timeout"
    finally:
        recu.close()
        proc.join(5 if donnees is not None else 0.5)
        if proc.is_alive():
            proc.kill()
            proc.join()
        if donnees is None and motif is None:
            motif = f"processus_pages_code_{proc.exitcode}"
        proc.close()
    return donnees, motif


#: Variables jamais transmises au processus qui analyse les fichiers déposés (contenu hostile) : clés, secrets,
#: URL de base (mot de passe éventuel). Le processus n'en a aucun besoin (revue de sécurité RS-14).
_SECRETS_ENV = re.compile(
    r"(CONTROLDONE_(MASTER_KEY|SECRET_KEY|DATABASE_URL|MCP_API_KEY|REFERENTIEL_SEL|IMAP_.*)|"
    r"ANTHROPIC_.*|STRIPE_.*|.*(SECRET|PASSWORD|PASSWD|TOKEN|API_KEY).*)",
    re.I,
)


def _environnement_sans_secrets() -> dict[str, str]:
    return {k: v for k, v in os.environ.items() if not _SECRETS_ENV.fullmatch(k)}


def _limiteur_memoire(memoire_mo: int):
    def _f():  # exécuté dans le processus enfant
        try:
            import resource

            octets = memoire_mo * 1024 * 1024
            resource.setrlimit(resource.RLIMIT_AS, (octets, octets))
        except (ImportError, ValueError, OSError):
            pass

    return _f


def _estimer_pages(contenu: bytes, mime: str) -> int:
    if mime == MIME_PDF:
        return max(1, len(re.findall(rb"/Type\s*/Page[^s]", contenu)))
    if mime == "image/tiff":
        try:
            from PIL import Image

            with Image.open(io.BytesIO(contenu)) as im:
                return max(1, int(getattr(im, "n_frames", 1)))
        except Exception:
            return 1
    return 1


def _page_illisible(numero: int, motif: str, **kw) -> PageText:
    return PageText(
        numero=numero, texte="", qualite=QualiteTexte.illisible, source="aucune", avertissements=[motif], **kw
    )


def extraire_pages_local(contenu: bytes, mime: str, options: OptionsPages | None = None) -> list[PageText]:
    """Extraction dans le processus courant (utilisée par le processus isolé)."""
    opts = options or OptionsPages()
    if mime == MIME_PDF:
        return _pages_pdf(contenu, opts)
    if mime.startswith("image/"):
        return _pages_image(contenu, opts)
    if mime == MIME_XLSX:
        return _pages_xlsx(contenu)
    if mime == MIME_ODS:
        return _pages_ods(contenu)
    if mime == MIME_CSV:
        return [_page_csv(contenu)]
    if mime == MIME_XML:
        return [_page_texte_brut(contenu, "xml")]
    if mime in (MIME_TEXTE, MIME_EML) or mime.startswith("text/"):
        return [_page_texte_brut(contenu, "texte")]
    return [_page_illisible(1, "type_non_lu")]


# --- PDF : texte natif ----------------------------------------------------------------------------------


def _couleur_blanche(c) -> bool:
    if c is None:
        return False
    if isinstance(c, (int, float)):
        c = (c,)
    try:
        vals = [float(x) for x in c]
    except (TypeError, ValueError):
        return False
    if len(vals) == 1:
        return vals[0] >= 0.95
    if len(vals) == 3:
        return all(v >= 0.95 for v in vals)
    if len(vals) == 4:
        return all(v <= 0.05 for v in vals)
    return False


def _fond_sombre(char: dict, rects: list[dict], images: list[dict]) -> bool:
    cx = (char["x0"] + char["x1"]) / 2
    cy = (char["top"] + char["bottom"]) / 2
    for r in rects:
        if (
            r.get("fill")
            and not _couleur_blanche(r.get("non_stroking_color"))
            and r["x0"] <= cx <= r["x1"]
            and r["top"] <= cy <= r["bottom"]
        ):
            return True
    return any(im["x0"] <= cx <= im["x1"] and im["top"] <= cy <= im["bottom"] for im in images)


def _caracteres_invisibles(pdfium_page) -> list[tuple[float, float, float, float]]:
    """Boîtes (x0, top, x1, bottom, en points depuis le haut) des caractères en mode de rendu invisible."""
    import pypdfium2.raw as raw

    boites: list[tuple[float, float, float, float]] = []
    try:
        tp = pdfium_page.get_textpage()
    except Exception:
        return boites
    hauteur = pdfium_page.get_height()
    try:
        n = raw.FPDFText_CountChars(tp.raw)
        import ctypes

        for i in range(n):
            obj = raw.FPDFText_GetTextObject(tp.raw, i)
            if not obj:
                continue
            if raw.FPDFTextObj_GetTextRenderMode(obj) != raw.FPDF_TEXTRENDERMODE_INVISIBLE:
                continue
            g, d, b, h = (ctypes.c_double(), ctypes.c_double(), ctypes.c_double(), ctypes.c_double())
            if raw.FPDFText_GetCharBox(tp.raw, i, g, d, b, h):
                # boîte étendue à la hauteur de la police (la boîte d'un « . » est minuscule)
                corps = float(raw.FPDFText_GetFontSize(tp.raw, i) or 0) or (h.value - b.value)
                ox, oy = ctypes.c_double(), ctypes.c_double()
                base = oy.value if raw.FPDFText_GetCharOrigin(tp.raw, i, ox, oy) else b.value
                haut = max(h.value, base + corps)
                bas = min(b.value, base - 0.3 * corps)
                boites.append((g.value - 1, hauteur - haut - 1, d.value + 1, hauteur - bas + 1))
    finally:
        tp.close()
    return boites


def _dans(boites, char) -> bool:
    cx = (char["x0"] + char["x1"]) / 2
    cy = (char["top"] + char["bottom"]) / 2
    return any(x0 <= cx <= x1 and t <= cy <= b for x0, t, x1, b in boites)


def _texte_natif(page, pdfium_page) -> tuple[list[Ligne], str, float, float]:
    """Lignes visibles, texte masqué, largeur, hauteur."""
    largeur, hauteur = float(page.width), float(page.height)
    rects = page.rects
    images = page.images
    invisibles = _caracteres_invisibles(pdfium_page) if pdfium_page is not None else []

    def cache(obj) -> bool:
        if obj.get("object_type") != "char":
            return False
        if (obj.get("size") or 0) < 2.0:
            return True
        if _couleur_blanche(obj.get("non_stroking_color")) and not _fond_sombre(obj, rects, images):
            return True
        return bool(invisibles) and _dans(invisibles, obj)

    visible = page.filter(lambda o: not cache(o))
    masque = page.filter(lambda o: o.get("object_type") == "char" and cache(o))
    try:
        texte_masque = masque.extract_text() or ""
    except Exception:
        texte_masque = ""
    mots_bruts = visible.extract_words(
        x_tolerance=1.5, y_tolerance=2.5, keep_blank_chars=False, use_text_flow=False, extra_attrs=["size"]
    )
    mots = []
    for w in mots_bruts:
        # Glyphe d'espace fine sans correspondance Unicode (police sous-ensemble) lu « \x00 » entre deux
        # chiffres : c'est un séparateur de milliers imprimé (« 1 405,44 »), restitué en espace fine.
        t = re.sub(r"(?<=\d)\x00(?=\d)", "\u202f", w["text"])
        if not t.strip():
            continue
        mots.append(
            Mot(
                texte=t,
                x0=w["x0"] / largeur,
                y0=w["top"] / hauteur,
                x1=w["x1"] / largeur,
                y1=w["bottom"] / hauteur,
                confiance=None,
                taille=round(float(w.get("size") or 0), 2),
            )
        )
    return construire_lignes(mots), texte_masque, largeur, hauteur


def _pages_pdf(contenu: bytes, opts: OptionsPages) -> list[PageText]:
    import pdfplumber
    import pypdfium2 as pdfium

    try:
        doc_pdfium = pdfium.PdfDocument(contenu)
    except Exception:
        doc_pdfium = None
    sortie: list[PageText] = []
    try:
        pdf = pdfplumber.open(io.BytesIO(contenu))
    except Exception:
        if doc_pdfium is None:
            return [_page_illisible(1, "pdf_illisible")]
        pdf = None
    try:
        n = len(pdf.pages) if pdf is not None else len(doc_pdfium)
        n = min(n, opts.max_pages)
        for i in range(n):
            numero = i + 1
            lignes: list[Ligne] = []
            texte_masque = ""
            largeur = hauteur = None
            avert: list[str] = []
            try:
                pp = doc_pdfium[i] if doc_pdfium is not None else None
            except Exception:
                pp = None
            if pdf is not None:
                try:
                    lignes, texte_masque, largeur, hauteur = _texte_natif(pdf.pages[i], pp)
                except Exception:
                    avert.append("texte_natif_illisible")
            texte = "\n".join(li.texte for li in lignes)
            nb = sum(1 for c in texte if not c.isspace())
            score = score_texte(texte) if nb else 0.0
            natif = PageText(
                numero=numero,
                texte=texte,
                lignes=lignes,
                qualite=QualiteTexte.natif,
                source="natif",
                score_natif=round(score, 4),
                largeur=largeur,
                hauteur=hauteur,
                texte_masque=texte_masque,
                avertissements=avert,
            )
            if nb >= MIN_CARACTERES_NATIF and score >= SEUIL_NATIF and not opts.forcer_ocr:
                sortie.append(natif)
                continue
            if not opts.ocr or pp is None or not ocr_disponible():
                if nb >= MIN_CARACTERES_NATIF and score >= SEUIL_NATIF_FAIBLE:
                    natif.qualite = QualiteTexte.natif_faible
                elif nb > 0 and score >= SEUIL_NATIF:
                    natif.qualite = QualiteTexte.natif_faible  # peu de texte, mais propre
                else:
                    natif.qualite = QualiteTexte.illisible
                    natif.avertissements.append("ocr_indisponible" if opts.ocr else "ocr_desactive")
                sortie.append(natif)
                continue
            try:
                echelle, reduite = echelle_rendu_ocr(*pp.get_size(), opts.dpi)
                image = pp.render(scale=echelle).to_pil()
                ocr = _ocr_image(image, opts, numero)
                if reduite:
                    ocr.avertissements.append("rendu_ocr_reduit")
            except Exception as e:  # rendu impossible
                ocr = _page_illisible(numero, f"rendu_impossible:{type(e).__name__}")
            ocr.largeur, ocr.hauteur, ocr.texte_masque = largeur, hauteur, texte_masque
            ocr.score_natif = natif.score_natif
            sortie.append(_meilleure(natif, ocr, nb))
    finally:
        if pdf is not None:
            pdf.close()
        if doc_pdfium is not None:
            doc_pdfium.close()
    return sortie or [_page_illisible(1, "pdf_sans_page")]


#: Plafond de pixels d'une page PDF rendue pour l'OCR (D-3606) : même borne que l'agrandissement des images
#: (REV2-03). Une page A2 à 300 dpi en fait 34,8 M, une A3 17,4 M : les pages ordinaires sont rendues à
#: ``OptionsPages.dpi`` sans changement. Au-delà (page de plusieurs mètres), la résolution baisse pour tenir dans
#: le plafond au lieu d'allouer des gigaoctets dans le processus isolé avant d'échouer sous ``RLIMIT_AS``.
MAX_PIXELS_RENDU_OCR = 40_000_000


def echelle_rendu_ocr(
    largeur_pt: float, hauteur_pt: float, dpi: int, max_pixels: int = MAX_PIXELS_RENDU_OCR
) -> tuple[float, bool]:
    """``(échelle pdfium, réduite)`` : ``dpi / 72``, abaissée si la page dépasserait ``max_pixels``."""
    w, h = max(1.0, float(largeur_pt)), max(1.0, float(hauteur_pt))
    echelle = dpi / 72
    if w * h * echelle * echelle <= max_pixels:
        return echelle, False
    return (max_pixels / (w * h)) ** 0.5 * 0.99, True  # marge pour les arrondis au pixel


def _meilleure(natif: PageText, ocr: PageText, nb_natif: int) -> PageText:
    """§7.3 : garder la meilleure des deux sources par page."""
    q_natif = (natif.score_natif or 0.0) if nb_natif >= MIN_CARACTERES_NATIF else 0.0
    q_ocr = (ocr.score_ocr or 0.0) * score_texte(ocr.texte) if ocr.texte.strip() else 0.0
    if q_natif >= SEUIL_NATIF_FAIBLE and q_natif >= q_ocr:
        natif.qualite = QualiteTexte.natif_faible
        natif.score_ocr = ocr.score_ocr
        return natif
    if (
        q_natif > 0
        and nb_natif > 0
        and natif.score_natif
        and natif.score_natif >= SEUIL_NATIF
        and ocr.qualite is QualiteTexte.illisible
    ):
        natif.qualite = QualiteTexte.natif_faible
        return natif
    return ocr


# --- OCR ------------------------------------------------------------------------------------------------


def _pour_tesseract(image):
    """Image passée à pytesseract, enregistrée en PGM/PPM (non compressé) au lieu de PNG.

    pytesseract écrit l'image dans un fichier temporaire au format ``image.format`` (PNG par défaut) : l'encodage
    PNG d'une page A4 à 300 dpi coûte ~1 s par appel Tesseract. Le PGM/PPM est sans perte comme le PNG : Tesseract
    lit exactement les mêmes pixels (texte identique), ~0,01 s d'écriture (D-1400)."""
    image.format = "PPM"  # mode « L » -> PGM (P5), « RGB » -> PPM (P6)
    return image


def _osd_rotation(image) -> int | None:
    """Rotation à appliquer (0/90/180/270) d'après l'OSD de Tesseract ; ``None`` sans verdict fiable.

    Réglage par défaut d'abord (plus fiable), puis avec un seuil de caractères abaissé (pages peu denses).
    Une page en paysage (scan pivoté) accepte un verdict 90/270 de moindre confiance."""
    import pytesseract

    paysage = image.size[0] > image.size[1]
    for config in ("--psm 0", "--psm 0 -c min_characters_to_try=10"):
        try:
            osd = pytesseract.image_to_osd(
                _pour_tesseract(image), config=config, output_type=pytesseract.Output.DICT, timeout=60
            )
        except Exception:
            continue
        rot = int(osd.get("rotate", 0)) % 360
        conf = float(osd.get("orientation_conf", 0.0))
        if rot in (90, 180, 270) and (conf >= 1.0 or (paysage and rot in (90, 270) and conf >= 0.5)):
            return rot
        if conf >= 1.0:
            return 0
    return None


def _angle_inclinaison(image) -> float:
    """Angle (degrés, sens trigonométrique) qui maximise la netteté du profil horizontal.

    PIL seul (D-1404 : plus de numpy) : image réduite à 1000 px, binarisée (encre = 255), profil = nombre de
    pixels d'encre par ligne, score = somme des carrés des différences entre lignes voisines (calcul exact en
    entiers)."""
    from PIL import Image, ImageStat

    petit = image.convert("L")
    w, h = petit.size
    facteur = 1000 / max(w, h)
    if facteur < 1:
        petit = petit.resize((max(1, int(w * facteur)), max(1, int(h * facteur))))
    seuil = min(200, int(ImageStat.Stat(petit).mean[0] * 0.8))
    binaire = petit.point([255 if v < seuil else 0 for v in range(256)])
    pw, ph = binaire.size
    if binaire.histogram()[255] < 0.002 * pw * ph:
        return 0.0

    def score(angle: float) -> int:
        donnees = binaire.rotate(angle, resample=Image.NEAREST, expand=False, fillcolor=0).tobytes()
        profil = [donnees.count(255, i, i + pw) for i in range(0, pw * ph, pw)]
        return sum((b - a) ** 2 for a, b in pairwise(profil))

    meilleur = max((a / 2 for a in range(-10, 11)), key=score)
    fin = max((meilleur + a / 10 for a in range(-5, 6)), key=score)
    return round(fin, 2)


def _ocr_brut(image, opts: OptionsPages) -> tuple[list[Mot], float]:
    import pytesseract

    data = pytesseract.image_to_data(
        _pour_tesseract(image),
        lang=opts.langues,
        config="--psm 3",
        output_type=pytesseract.Output.DICT,
        timeout=max(30, int(opts.timeout_par_page_s)),
    )
    w, h = image.size
    mots: list[Mot] = []
    poids = 0
    somme = 0.0
    for i, t in enumerate(data["text"]):
        t = (t or "").strip()
        try:
            conf = float(data["conf"][i])
        except (TypeError, ValueError):
            conf = -1.0
        if not t or conf < 0:
            continue
        x, y, ww, hh = data["left"][i], data["top"][i], data["width"][i], data["height"][i]
        mots.append(
            Mot(
                texte=t,
                x0=x / w,
                y0=y / h,
                x1=(x + ww) / w,
                y1=(y + hh) / h,
                confiance=round(conf / 100, 4),
                taille=round(hh / h, 5),
            )
        )
        somme += conf / 100 * len(t)
        poids += len(t)
    return mots, (somme / poids if poids else 0.0)


#: Réglages du prétraitement OCR (D-2601 à D-2606). Un dictionnaire pour que ``scripts/mesure_ocr.py`` puisse
#: comparer des variantes ; le pipeline n'en modifie jamais les valeurs.
REGLAGES_OCR: dict[str, object] = {
    "gris_max": True,  # gris = max(R, V, B) : encre colorée claire atténuée (D-2601)
    "traits": True,  # traits parasites traversant la page effacés (D-2602)
    "median": True,  # filtre médian 3 × 3 si bruit impulsionnel (D-2602)
    "contraste": True,  # étirement d'une page pâle (D-2603)
    "deux_pages": True,  # feuille « deux pages par feuille » lue moitié par moitié (D-2604)
    "seuil_reessai_orientation": 0.40,  # qualité de lecture sous laquelle les autres orientations sont essayées (D-2605)
}


def _qualite_ocr(mots: list[Mot], score: float) -> float:
    return score * score_texte(" ".join(m.texte for m in mots))


def _pretraiter(image) -> tuple[object, list[str]]:
    """Image ``L`` nettoyée avant orientation et OCR (``ingest.pretraitement``) et avertissements de traçabilité."""
    from . import pretraitement as pt

    r = REGLAGES_OCR
    gris = pt.niveaux_de_gris(image) if r["gris_max"] else image.convert("L")
    notes: list[str] = []
    if r["traits"]:
        gris, n = pt.retirer_traits(gris)
        if n:
            notes.append(f"pretraitement:traits_effaces:{n}")
    if r["median"]:  # avant l'étirement : la détection de télécopie lit les demi-teintes d'origine
        gris, filtre = pt.debruiter(gris)
        if filtre:
            notes.append("pretraitement:median")
    if r["contraste"]:
        gris, etire = pt.etirer_contraste(gris)
        if etire:
            notes.append("pretraitement:contraste_etire")
    return gris, notes


def _ocr_oriente(
    image, opts: OptionsPages, rotation: int | None
) -> tuple[list[Mot], float, int, float, object]:
    """Désinclinaison, OCR, puis essai des autres orientations si la lecture est mauvaise.

    ``rotation`` : verdict OSD déjà appliqué (``None`` : pas de verdict). Autres orientations essayées quand l'OSD
    n'a pas de verdict et que la lecture est moyenne (confiance < 0,6 ou moins de 5 mots), ou — même avec un
    verdict, D-2113/D-2605 — quand la qualité de lecture (confiance × plausibilité du texte, ``score_texte``) est
    sous ``seuil_reessai_orientation`` : OSD confiant mais faux. Une page lue à l'envers garde souvent une
    confiance Tesseract de 0,35–0,45 ; sa qualité tombe à 0,25–0,33 (texte implausible), celle d'une page droite
    dépasse 0,8. On garde la lecture de meilleure qualité (avance d'au moins 0,1)."""
    from PIL import Image

    angle = 0.0
    try:
        angle = _angle_inclinaison(image)
    except Exception:
        angle = 0.0
    if abs(angle) >= 0.3:
        image = image.rotate(angle, resample=Image.BICUBIC, expand=True, fillcolor=255)
    mots, score = _ocr_brut(image, opts)
    deja = rotation or 0
    seuil = float(REGLAGES_OCR["seuil_reessai_orientation"] or 0.0)
    q_meilleur = _qualite_ocr(mots, score)
    if (rotation is None and (score < 0.6 or len(mots) < 5)) or q_meilleur < seuil:
        meilleurs = (mots, score, deja, image)
        for cible in (0, 180, 90, 270):  # 0 et 180 d'abord : Tesseract (psm 3) relit seul une page à 90°
            if cible == deja:
                continue
            im2 = image.rotate(-((cible - deja) % 360), expand=True, fillcolor=255)
            m2, s2 = _ocr_brut(im2, opts)
            q2 = _qualite_ocr(m2, s2)
            if q2 > q_meilleur + 0.1:
                meilleurs, q_meilleur = (m2, s2, cible, im2), q2
        mots, score, deja, image = meilleurs
    return mots, score, deja, angle, image


def _rotation_deux_pages(rotation: int | None) -> int | None:
    """Feuille « deux pages » détectée avant l'OSD : seul un retournement (180°) est appliqué."""
    return rotation if rotation in (0, 180) else None


def _ocr_image(image, opts: OptionsPages, numero: int) -> PageText:
    from . import pretraitement as pt

    # Tesseract multi-fils se dégrade fortement sous charge (plusieurs pages en parallèle) : un fil par OCR.
    os.environ.setdefault("OMP_THREAD_LIMIT", "1")

    image, notes = _pretraiter(image)
    coupe = pt.coupure_deux_pages(image) if REGLAGES_OCR["deux_pages"] else None
    rotation = _osd_rotation(image)
    if coupe is not None:
        rotation = _rotation_deux_pages(rotation)
    if rotation:
        image = image.rotate(-rotation, expand=True, fillcolor=255)
        if REGLAGES_OCR["deux_pages"]:
            coupe = pt.coupure_deux_pages(image)
    if coupe is not None:
        double = _ocr_deux_pages(image, coupe, opts, numero, rotation or 0, notes)
        if moities_separees(double.lignes, coupe / image.size[0]):
            return double
        # D-2902 : du texte touche la coupure (tableau en paysage dont deux colonnes encadrent la bande centrale,
        # pied de page centré) : une seule page, relue entière
        notes = [*notes, "deux_pages_ecarte"]
    mots, score, rot_finale, angle, _ = _ocr_oriente(image, opts, rotation)
    return _page_ocr(numero, mots, score, rot_finale, angle, notes)


def _page_ocr(
    numero: int,
    mots: list[Mot],
    score: float,
    rotation: int,
    angle: float,
    notes: list[str],
    lignes: list[Ligne] | None = None,
) -> PageText:
    lignes = construire_lignes(mots) if lignes is None else lignes
    texte = "\n".join(li.texte for li in lignes)
    qualite = QualiteTexte.ocr
    if not texte.strip() or score < SEUIL_ILLISIBLE_OCR or score_texte(texte) < 0.4:
        qualite = QualiteTexte.illisible
    return PageText(
        numero=numero,
        texte=texte,
        lignes=lignes,
        qualite=qualite,
        source="ocr",
        score_ocr=round(score, 4),
        rotation=rotation,
        desinclinaison=angle,
        avertissements=notes,
    )


#: Distance minimale (fraction de la largeur) entre le texte de chaque moitié et la coupure d'une vraie feuille
#: « deux pages » : les marges des deux pages réduites l'encadrent (≥ 3 % de chaque côté sur le banc, D-2902).
MARGE_COUPURE_DEUX_PAGES = 0.012


def moities_separees(lignes: Sequence[Ligne], coupe: float) -> bool:
    """Les lignes lues de part et d'autre de la coupure ``coupe`` (0–1) restent à distance de celle-ci : une
    ligne qui la touche (moins de ``MARGE_COUPURE_DEUX_PAGES``) appartient à une page unique coupée à tort
    (une colonne vide au milieu d'un tableau en paysage n'est pas une gouttière, D-2902)."""
    for li in lignes:
        mots = [m for m in li.mots if sum(c.isalnum() for c in m.texte) >= 2]
        if not mots:
            continue
        x0, x1 = min(m.x0 for m in mots), max(m.x1 for m in mots)
        if x1 <= coupe + 0.002 and coupe - x1 < MARGE_COUPURE_DEUX_PAGES:
            return False
        if x0 >= coupe - 0.002 and x0 - coupe < MARGE_COUPURE_DEUX_PAGES:
            return False
    return True


def _ocr_deux_pages(
    image, coupe: int, opts: OptionsPages, numero: int, rotation: int, notes: list[str]
) -> PageText:
    """Feuille « deux pages par feuille » (D-2604) : chaque moitié est lue seule (désinclinaison, OCR, orientation),
    puis les lignes de la moitié gauche précèdent celles de la moitié droite.

    Convention de numérotation : la feuille reste **une** page physique (même ``numero``, mêmes citations « page
    n ») ; les boîtes des mots sont normalisées sur la feuille entière (moitié droite : x ≥ abscisse de coupe),
    si bien qu'une preuve citée pointe au bon endroit de la feuille. Les lignes ne mêlent jamais les deux
    moitiés. Avertissement ``deux_pages_par_feuille:<coupe relative>``."""
    w, h = image.size
    mots_tous: list[Mot] = []
    lignes: list[Ligne] = []
    angles: list[float] = []
    somme = poids = 0.0
    for x0, x1 in ((0, coupe), (coupe, w)):
        mots, score, _rot, angle, _im = _ocr_oriente(image.crop((x0, 0, x1, h)), opts, 0)
        angles.append(angle)
        # coordonnées 0–1 de la moitié -> feuille (abscisses décalées et réduites, ordonnées inchangées)
        places = [replace(m, x0=(x0 + m.x0 * (x1 - x0)) / w, x1=(x0 + m.x1 * (x1 - x0)) / w) for m in mots]
        mots_tous.extend(places)
        lignes.extend(construire_lignes(places))
        n = sum(len(m.texte) for m in mots)
        somme += score * n
        poids += n
    score = somme / poids if poids else 0.0
    notes = [*notes, f"deux_pages_par_feuille:{coupe / w:.3f}"]
    return _page_ocr(numero, mots_tous, score, rotation, angles[0] if angles else 0.0, notes, lignes=lignes)


#: Taille maximale (pixels) d'une image agrandie avant l'OCR : une page A3 à 300 dpi en fait 17,4 M.
MAX_PIXELS_AGRANDISSEMENT = 40_000_000


def _pages_image(contenu: bytes, opts: OptionsPages) -> list[PageText]:
    from PIL import Image, ImageSequence

    sortie: list[PageText] = []
    try:
        im = Image.open(io.BytesIO(contenu))
        # Cadres décodés un par un (REV2-03) : copier tous les cadres avant de couper à ``max_pages`` gardait en
        # mémoire chaque page d'un TIFF ou d'un GIF à cadres multiples.
        cadres = itertools.islice(ImageSequence.Iterator(im), opts.max_pages)
    except Exception:
        return [_page_illisible(1, "image_illisible")]
    for i in itertools.count(1):
        try:
            cadre = next(cadres).copy()
        except StopIteration:
            break
        except Exception:
            if not sortie:
                return [_page_illisible(1, "image_illisible")]
            break
        if cadre.mode in ("RGBA", "LA", "P"):
            fond = Image.new("RGB", cadre.size, (255, 255, 255))
            cadre = cadre.convert("RGBA")
            fond.paste(cadre, mask=cadre.split()[-1])
            cadre = fond
        if not opts.ocr or not ocr_disponible():
            sortie.append(
                _page_illisible(
                    i,
                    "ocr_indisponible" if opts.ocr else "ocr_desactive",
                    largeur=float(cadre.size[0]),
                    hauteur=float(cadre.size[1]),
                )
            )
            continue
        dpi = cadre.info.get("dpi", (0, 0))[0] or 0
        if dpi and dpi < 150:  # agrandir une image basse résolution avant l'OCR
            # sans dépasser MAX_PIXELS_AGRANDISSEMENT (REV2-03 : une image de 80 Mpx à 72 dpi en demandait 720)
            f = min(
                3.0, 300 / dpi, (MAX_PIXELS_AGRANDISSEMENT / max(1, cadre.size[0] * cadre.size[1])) ** 0.5
            )
            if f > 1:
                cadre = cadre.resize((int(cadre.size[0] * f), int(cadre.size[1] * f)))
        try:
            p = _ocr_image(cadre, opts, i)
        except Exception as e:
            p = _page_illisible(i, f"ocr_echec:{type(e).__name__}")
        p.largeur, p.hauteur = float(cadre.size[0]), float(cadre.size[1])
        sortie.append(p)
    return sortie or [_page_illisible(1, "image_vide")]


# --- tableurs, CSV, XML, texte ----------------------------------------------------------------------


def _decimales_format(fmt: str | None) -> int | None:
    if not fmt or fmt == "General":
        return None
    m = re.search(r"0\.(0+)", fmt)
    if m:
        return len(m.group(1))
    if re.search(r"(^|[^.])0(?![.0])", fmt) or "#,##0" in fmt:
        return 0
    return None


def _cellule_texte(v, fmt: str | None = None) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VRAI" if v else "FAUX"
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0) else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float, Decimal)):
        try:
            d = Decimal(repr(v)) if isinstance(v, float) else Decimal(v)
        except InvalidOperation:
            return str(v)
        n = _decimales_format(fmt)
        if n is not None and "%" not in (fmt or ""):
            d = d.quantize(Decimal(1).scaleb(-n)) if n else d.quantize(Decimal(1))
        s = format(d, "f")
        if isinstance(v, float) and n is None and s.endswith(".0"):
            s = s[:-2]
        return s
    return str(v).replace("\r", " ").replace("\n", " ").strip()


def _page_grille(
    numero: int,
    lignes_cellules: list[list[str]],
    *,
    source: str,
    feuille: str | None,
    texte_brut: str | None = None,
) -> PageText:
    """Page à partir d'une grille de cellules : texte stable (cellules non vides jointes par `` | ``),
    mots positionnés sur une grille régulière."""
    lignes_cellules = [r for r in lignes_cellules]
    while lignes_cellules and not any(c for c in lignes_cellules[-1]):
        lignes_cellules.pop()
    nrows = max(1, len(lignes_cellules))
    ncols = max([len(r) for r in lignes_cellules] + [1])
    lignes: list[Ligne] = []
    textes: list[str] = []
    n_mots = 0
    for r, cellules in enumerate(lignes_cellules):
        valeurs = [(c, val) for c, val in enumerate(cellules) if val]
        if not valeurs:
            continue
        texte_ligne = " | ".join(v for _c, v in valeurs)
        textes.append(texte_ligne)
        if len(lignes) < MAX_LIGNES_POSITIONNEES and n_mots < MAX_MOTS_POSITIONNES:
            mots = tuple(
                Mot(texte=val, x0=c / ncols, y0=r / nrows, x1=(c + 1) / ncols, y1=(r + 1) / nrows)
                for c, val in valeurs
            )
            n_mots += len(mots)
            lignes.append(Ligne(texte=texte_ligne, mots=mots))
    texte = texte_brut if texte_brut is not None else "\n".join(textes)
    avert = ["texte_positionne_tronque"] if len(lignes) < len(textes) else []
    return PageText(
        numero=numero,
        texte=texte,
        lignes=lignes,
        qualite=QualiteTexte.natif,
        source=source,
        score_natif=1.0,
        feuille=feuille,
        avertissements=avert,
    )


def _pages_xlsx(contenu: bytes) -> list[PageText]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    except Exception:
        return [_page_illisible(1, "tableur_illisible")]
    sortie = []
    try:
        for i, ws in enumerate(wb.worksheets, start=1):
            grille: list[list[str]] = []
            cellules = 0
            tronque = False
            # dimensions annoncées non fiables (« A1:XFD1048576 ») : colonnes et cellules lues bornées (D-1604)
            for row in ws.iter_rows(max_col=MAX_COLONNES_TABLEUR):
                grille.append(
                    [
                        _cellule_texte(getattr(c, "value", None), getattr(c, "number_format", None))
                        for c in row
                    ]
                )
                cellules += max(1, len(row))
                if len(grille) > 100_000 or cellules > MAX_CELLULES_FEUILLE:
                    tronque = True
                    break
            page = _page_grille(i, grille, source="tableur", feuille=ws.title)
            if tronque:
                page.avertissements.append("tableur_tronque")
            sortie.append(page)
    finally:
        wb.close()
    return sortie or [_page_illisible(1, "tableur_vide")]


_NS_ODS = {
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}


def _pages_ods(contenu: bytes) -> list[PageText]:
    from lxml import etree

    try:
        with zipfile.ZipFile(io.BytesIO(contenu)) as z:
            xml = z.read("content.xml")
        parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
        racine = etree.fromstring(xml, parser)
    except Exception:
        return [_page_illisible(1, "tableur_illisible")]
    t, o = "{{{}}}".format(_NS_ODS["table"]), "{{{}}}".format(_NS_ODS["office"])
    sortie = []
    for i, table in enumerate(racine.iter(f"{t}table"), start=1):
        grille: list[list[str]] = []
        for row in table.iter(f"{t}table-row"):
            cellules: list[str] = []
            for cell in row:
                if cell.tag not in (f"{t}table-cell", f"{t}covered-table-cell"):
                    continue
                rep = min(int(cell.get(f"{t}number-columns-repeated", "1")), 1000)
                vt = cell.get(f"{o}value-type")
                if vt in ("float", "currency", "percentage"):
                    val = cell.get(f"{o}value", "")
                elif vt == "date":
                    val = (cell.get(f"{o}date-value") or "")[:10]
                else:
                    val = " ".join(
                        "".join(p.itertext()) for p in cell.iter("{{{}}}p".format(_NS_ODS["text"]))
                    ).strip()
                cellules.extend([val] * (rep if val else min(rep, 50)))
            while cellules and not cellules[-1]:
                cellules.pop()
            rep_l = min(int(row.get(f"{t}number-rows-repeated", "1")), 1000 if cellules else 1)
            grille.extend([cellules] * rep_l)
        sortie.append(_page_grille(i, grille, source="tableur", feuille=table.get(f"{t}name")))
    return sortie or [_page_illisible(1, "tableur_vide")]


def _page_csv(contenu: bytes) -> PageText:
    import csv

    texte = decoder_texte(contenu)
    if texte is None:
        return _page_illisible(1, "csv_illisible")
    texte = texte.replace("\r\n", "\n").replace("\r", "\n")
    lignes = [li for li in texte.split("\n")]
    try:
        dialecte = csv.Sniffer().sniff("\n".join(lignes[:20]), delimiters=";,\t|")
        grille = [[c.strip() for c in r] for r in csv.reader(lignes, dialecte)]
    except csv.Error:
        grille = [[li] for li in lignes]
    p = _page_grille(1, grille, source="csv", feuille=None, texte_brut=texte.rstrip("\n"))
    return p


def _page_texte_brut(contenu: bytes, source: str) -> PageText:
    texte = decoder_texte(contenu)
    if texte is None:
        try:
            texte = contenu.decode("utf-8", "replace")
        except Exception:
            return _page_illisible(1, "texte_illisible")
    texte = texte.lstrip("﻿").replace("\r\n", "\n").replace("\r", "\n")
    lignes_txt = texte.split("\n")
    n = max(1, len(lignes_txt))
    largeur_max = max(1, max((len(x) for x in lignes_txt), default=1))  # une fois (D-1401 : était O(n²))
    lignes = []
    n_mots = 0
    tronque = False
    for r, li in enumerate(lignes_txt):
        if len(lignes) >= MAX_LIGNES_POSITIONNEES or n_mots >= MAX_MOTS_POSITIONNES:
            tronque = True  # texte complet conservé ; positions bornées (D-1604)
            break
        mots = []
        for m in re.finditer(r"\S+", li):
            mots.append(
                Mot(
                    texte=m.group(),
                    x0=m.start() / largeur_max,
                    y0=r / n,
                    x1=min(1.0, m.end() / largeur_max),
                    y1=(r + 1) / n,
                )
            )
            if n_mots + len(mots) >= MAX_MOTS_POSITIONNES:
                tronque = True
                break
        n_mots += len(mots)
        if mots:
            lignes.append(Ligne(texte=li.strip(), mots=tuple(mots)))
    return PageText(
        numero=1,
        texte=texte,
        lignes=lignes,
        qualite=QualiteTexte.natif,
        source=source,
        score_natif=1.0,
        avertissements=["texte_positionne_tronque"] if tronque else [],
    )


__all__ += ["Ligne", "Mot", "PageText"]

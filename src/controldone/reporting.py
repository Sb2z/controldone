"""Excel reporter for v4."""
from __future__ import annotations

import os
import re
from typing import Iterable, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from controldone.models import Bundle, CheckResult, Document
from controldone.validation import blocking_reasons


FILL_HEADER = PatternFill("solid", fgColor="1F4E79")
FILL_SECTION = PatternFill("solid", fgColor="DDEBF7")
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_WARN = PatternFill("solid", fgColor="FFEB9C")
FILL_NOK = PatternFill("solid", fgColor="FFC7CE")
FILL_SKIP = PatternFill("solid", fgColor="D9EAD3")
FONT_HEADER = Font(bold=True, color="FFFFFF", name="Arial", size=10)
FONT_BOLD = Font(bold=True, name="Arial", size=10)
FONT_NORMAL = Font(name="Arial", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_report_v4(bundles: List[Bundle], out_path: str,
                    errors: List[Tuple[str, str]] | None = None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthese"
    _write_synthesis(ws, bundles)

    existing = {"Synthese"}
    for idx, bundle in enumerate(bundles, start=1):
        name = _safe_tab(bundle.folder_id or f"DOSSIER_{idx}", existing)
        _write_detail(wb.create_sheet(name), bundle)

    _write_warnings(wb.create_sheet("OK_A_Controler"), bundles)
    _write_issues(wb.create_sheet("KO_Bloquants"), bundles)
    if errors:
        _write_errors(wb.create_sheet("Erreurs"), errors)

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    wb.save(out_path)


def _write_synthesis(ws, bundles: List[Bundle]) -> None:
    headers = [
        "N facture prestation", "AWB/LTA", "MRN", "Facture marchandise",
        "TVA facture", "TVA declaration", "Statut", "Raisons",
        "Montant facture", "Devise facture", "Montant declaration",
        "Devise declaration", "HS facture", "HS declaration", "Sources",
    ]
    ws.append(headers)
    _format_header(ws, len(headers))
    for b in bundles:
        inv = b.invoices[0] if b.invoices else None
        dec = b.declarations[0] if b.declarations else None
        ws.append([
            b.ref_prestation_no,
            b.ref_awb_lta,
            _join(d.mrn for d in b.declarations) or b.ref_mrn,
            _join(d.invoice_number for d in b.invoices),
            _join_vats(d.vat_buyer for d in b.invoices),
            _join_vats((d.importer.vat or d.buyer.vat) for d in b.declarations),
            b.status,
            blocking_reasons(b.checks),
            _sum_amount(d.total_amount for d in b.invoices),
            b.primary_currency or "",
            _sum_amount(d.total_amount for d in b.declarations),
            dec.currency if dec else "",
            _join_hs(b.invoices),
            _join_hs(b.declarations),
            _join(os.path.basename(x) for x in b.source_files),
        ])
        _format_row(ws, ws.max_row, b.status)
    widths = [22, 18, 24, 24, 18, 18, 16, 60, 16, 14, 18, 14, 28, 28, 55]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write_detail(ws, b: Bundle) -> None:
    ws.append(["Reference dossier", b.folder_id])
    ws.append(["N facture prestation", b.ref_prestation_no])
    ws.append(["AWB/LTA", b.ref_awb_lta])
    ws.append(["MRN", b.ref_mrn])
    ws.append(["Sources", _join(os.path.basename(x) for x in b.source_files)])
    ws.append([])

    ws.append(["Documents"])
    _section(ws, ws.max_row, 8)
    ws.append(["Type", "Format", "Fichier", "Pages", "N facture", "MRN", "AWB", "TVA"])
    _format_header(ws, 8, ws.max_row)
    for d in b.invoices + b.declarations + ([b.prestation] if b.prestation else []):
        ws.append([
            d.kind, d.format, os.path.basename(d.source_file or ""),
            ",".join(map(str, d.source_pages)), d.invoice_number or "",
            d.mrn or "", d.awb or "", d.vat_buyer or d.importer.vat or "",
        ])

    ws.append([])
    ws.append(["Controles"])
    _section(ws, ws.max_row, 8)
    ws.append(["Section", "Controle", "Attendu", "Obtenu", "Statut", "Severite", "Commentaire", "Code"])
    _format_header(ws, 8, ws.max_row)
    for c in b.checks:
        ws.append([c.section, c.label, c.expected, c.got, c.status, c.severity, c.comment, c.code])
        _format_check_row(ws, ws.max_row, c.status)

    ws.append([])
    ws.append(["Articles facture"])
    _section(ws, ws.max_row, 8)
    ws.append(["N", "HS", "Description", "Origine", "Montant", "Devise", "Poids net", "Poids brut"])
    _format_header(ws, 8, ws.max_row)
    for d in b.invoices:
        for ln in d.lines:
            ws.append([ln.item_no, ln.hs_code, ln.description, ln.origin_country, ln.amount, ln.currency, ln.net_weight_kg, ln.gross_weight_kg])

    ws.append([])
    ws.append(["Articles declaration"])
    _section(ws, ws.max_row, 8)
    ws.append(["N", "HS", "Description", "Origine", "Montant", "Devise", "Poids net", "Poids brut", "Preference"])
    _format_header(ws, 9, ws.max_row)
    for d in b.declarations:
        for ln in d.lines:
            ws.append([ln.item_no, ln.hs_code, ln.description, ln.origin_country, ln.amount, ln.currency, ln.net_weight_kg, ln.gross_weight_kg, ln.preference])

    for col in range(1, 10):
        ws.column_dimensions[get_column_letter(col)].width = [18, 24, 36, 18, 18, 12, 14, 14, 14][col - 1]
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER


def _write_issues(ws, bundles: List[Bundle]) -> None:
    ws.append(["Reference", "Statut", "Section", "Controle", "Attendu", "Obtenu", "Commentaire", "Sources"])
    _format_header(ws, 8)
    for b in bundles:
        for c in b.checks:
            if c.is_blocking:
                ws.append([b.folder_id, b.status, c.section, c.label, c.expected, c.got, c.comment, _join(os.path.basename(x) for x in b.source_files)])
                _format_check_row(ws, ws.max_row, c.status)
    for i, w in enumerate([22, 16, 18, 34, 24, 24, 55, 55], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_warnings(ws, bundles: List[Bundle]) -> None:
    ws.append(["Reference", "Statut", "Section", "Controle", "Attendu", "Obtenu", "Commentaire", "Sources"])
    _format_header(ws, 8)
    for b in bundles:
        for c in b.checks:
            if c.status == "WARN":
                ws.append([b.folder_id, b.status, c.section, c.label, c.expected, c.got, c.comment, _join(os.path.basename(x) for x in b.source_files)])
                _format_check_row(ws, ws.max_row, c.status)
    for i, w in enumerate([22, 18, 18, 34, 24, 32, 65, 55], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_errors(ws, errors: List[Tuple[str, str]]) -> None:
    ws.append(["Fichier", "Erreur"])
    _format_header(ws, 2)
    for path, err in errors:
        ws.append([os.path.basename(path), err])
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 90


def _format_header(ws, cols: int, row: int = 1) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _format_row(ws, row: int, status: str) -> None:
    fill = FILL_OK if status == "OK" else FILL_WARN if status in {"A_VERIFIER", "OK_A_CONTROLER"} else FILL_NOK
    for cell in ws[row]:
        cell.fill = fill
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = FONT_NORMAL


def _format_check_row(ws, row: int, status: str) -> None:
    fill = {"OK": FILL_OK, "WARN": FILL_WARN, "NOK": FILL_NOK, "SKIP": FILL_SKIP}.get(status, FILL_SKIP)
    for cell in ws[row]:
        cell.fill = fill
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = FONT_NORMAL


def _section(ws, row: int, cols: int) -> None:
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = FILL_SECTION
        cell.font = FONT_BOLD
        cell.border = BORDER


def _safe_tab(name: str, existing: set) -> str:
    name = re.sub(r"[:\\/?*\[\]]", "_", str(name or "DOSSIER"))[:31]
    base = name or "DOSSIER"
    i = 2
    while name in existing:
        suffix = f"_{i}"
        name = (base[:31 - len(suffix)] + suffix)
        i += 1
    existing.add(name)
    return name


def _join(values: Iterable[object]) -> str:
    vals = [str(v) for v in values if v not in (None, "")]
    return ", ".join(vals)


def _join_hs(docs: Iterable[Document]) -> str:
    vals = []
    for d in docs:
        vals.extend(d.hs_codes)
    return _join(dict.fromkeys(vals))


def _join_vats(values: Iterable[object]) -> str:
    return _join(dict.fromkeys(v for v in values if v not in (None, "")))


def _sum_amount(values: Iterable[float | None]) -> str:
    nums = [v for v in values if v is not None]
    return f"{sum(nums):.2f}" if nums else ""

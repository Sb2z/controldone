"""Corrections du moteur issues de l'audit final (D-12xx) — données fictives."""

from __future__ import annotations

import pytest
from test_pipeline import PROFIL, _composants, lot  # noqa: F401  (fixture réutilisée)

from controldone.controls.context import ControlContext
from controldone.findings_io import statut_global_depuis_resultats
from controldone.model import Outcome, RaisonCode
from controldone.model.enums import StatutGlobal
from controldone.pipeline import OptionsPipeline, traiter_lot


def test_f4_aucun_resultat_n_est_jamais_conforme():
    assert statut_global_depuis_resultats([]) is StatutGlobal.a_verifier


def test_f4_contexte_en_erreur_donne_a_verifier(lot, monkeypatch):  # noqa: F811
    def casse(*a, **k):
        raise RuntimeError("panne fictive")

    monkeypatch.setattr(ControlContext, "construire", classmethod(lambda cls, *a, **k: casse()))
    res = traiter_lot(lot / "docs", PROFIL, [], options=OptionsPipeline(seed=1), composants=_composants())
    r = res[0]
    assert r.statut_global is StatutGlobal.a_verifier
    assert r.findings.statut_global == "a_verifier"
    assert [(x.outcome, x.raison_code) for x in r.resultats] == [(Outcome.non_verifiable, RaisonCode.erreur_interne)]


def test_f10_memo_ne_rejoue_pas_un_sous_ensemble_de_controles(lot):  # noqa: F811
    memo: dict = {}
    comp = _composants()
    a = traiter_lot(lot / "docs", PROFIL, [], options=OptionsPipeline(seed=1, memo=memo, controles=["A1"]),
                    composants=comp)
    b = traiter_lot(lot / "docs", PROFIL, [], options=OptionsPipeline(seed=1, memo=memo), composants=comp)
    assert {x.controle_id for x in a[0].resultats} == {"A1"}
    assert {x.controle_id for x in b[0].resultats} > {"A1", "B1"}


def test_f10_memo_depend_des_grilles_et_du_profil(lot):  # noqa: F811
    from controldone.pipeline import empreinte_contexte_controles
    from controldone.referentiel_io import profil_depuis_dict
    from controldone.taux_reference import TableTauxReference

    p1 = profil_depuis_dict(PROFIL)
    p2 = profil_depuis_dict({**PROFIL, "transitaires": []})
    t = TableTauxReference({})
    e1 = empreinte_contexte_controles(p1, [], controles=None, taux_reference=t)
    assert e1 == empreinte_contexte_controles(p1, [], controles=None, taux_reference=t)
    assert e1 != empreinte_contexte_controles(p2, [], controles=None, taux_reference=t)
    assert e1 != empreinte_contexte_controles(p1, [], controles=["A1"], taux_reference=t)
    from datetime import date
    from decimal import Decimal
    t2 = TableTauxReference({"USD": {date(2026, 1, 2): Decimal("1.1")}})
    assert e1 != empreinte_contexte_controles(p1, [], controles=None, taux_reference=t2)


if __name__ == "__main__":  # pragma: no cover
    pytest.main([__file__])


# --- F1, F8 : totaux du rapport ; F5 : JSON publiés ; F14 : tableur sans float ---------------------------

def _rd(lot):  # noqa: F811
    return traiter_lot(lot / "docs", PROFIL, [], options=OptionsPipeline(seed=1, annee=2026), composants=_composants())[0]


def _resultat(rd, cid, niveau, montant, nature, *, docs=(), details=None, unite="u"):
    from decimal import Decimal

    from controldone.model import Constat, Niveau, ResultatControle
    from controldone.model.enums import NatureMontant

    niv = Niveau(niveau)
    raisons = [RaisonCode.controle_signal_seulement] if niv is Niveau.a_verifier else []
    c = Constat(controle_id=cid, niveau=niv, raisons=raisons, libelle=f"constat fictif {cid}",
                montant_en_jeu=Decimal(montant), nature_montant=NatureMontant(nature), documents_concernes=list(docs))
    return ResultatControle(controle_id=cid, unite=unite, dossier_id=rd.dossier.id, dossier_version=1,
                            outcome=niv.outcome, constat=c, details=details or {}, documents_concernes=list(docs))


def test_f1_e6_remplace_le_constat_d_origine_dans_les_totaux(lot):  # noqa: F811
    from controldone.findings_io import construire_findings
    from controldone.rapport.vue import construire_vue

    rd = _rd(lot)
    rd.resultats = [r for r in rd.resultats if r.constat is None]
    d3 = _resultat(rd, "D3", "ecart_certain", "29.56", "recouvrable")
    e6 = _resultat(rd, "E6", "a_verifier", "10.63", "recouvrable",
                   details={"remplace_constat_id": d3.constat.id})
    rd.resultats += [d3, e6]
    v = construire_vue([rd], rd.profil)
    dv = v.dossiers[0]
    assert dv.recouvrable_certain == 0 and str(dv.recouvrable_a_verifier) == "10.63"
    assert v.recouvrable_certain.startswith("0,00") and v.recouvrable_a_verifier.startswith("10,63")
    # le constat d'origine reste listé, avec son montant
    assert {c.controle_id for c in dv.constats} == {"D3", "E6"}
    f = construire_findings(rd.dossier, rd.documents.values(), rd.resultats, rd.execution)
    e6f = next(c for c in f.constats if c.controle_id == "E6")
    assert e6f.remplace_constat_id == d3.constat.id


def test_f8_f5_et_a4_meme_ecart_comptes_une_fois(lot):  # noqa: F811
    from controldone.rapport.vue import construire_vue

    rd = _rd(lot)
    rd.resultats = [r for r in rd.resultats if r.constat is None]
    fc = next(i for i, d in rd.documents.items() if d.type.value == "facture_commerciale")
    rd.resultats += [
        _resultat(rd, "A4", "ecart_certain", "3497.65", "ecart_documentaire", docs=[fc]),
        _resultat(rd, "F5", "a_verifier", "3497.65", "ecart_documentaire", docs=[fc]),
    ]
    v = construire_vue([rd], rd.profil)
    assert v.ecarts_documentaires_nb == 1
    assert v.ecarts_documentaires.replace(" ", " ").replace("\xa0", " ").startswith("3 497,65")


def test_f5_libelle_bloque_absent_des_json_publies(lot):  # noqa: F811
    import json

    from controldone.rapport.export import findings_lot_json

    rd = _rd(lot)
    c0 = rd.findings.constats[0]
    bloque = c0.model_copy(update={"libelle": "Le droit dû est de 9,54 EUR", "motif_blocage": "formulation_interdite"})
    rd.findings = rd.findings.model_copy(update={"constats": [bloque, *rd.findings.constats[1:]]})
    texte = json.dumps(findings_lot_json([rd]), ensure_ascii=False)
    assert "droit dû" not in texte and "Libellé retenu pour relecture" in texte


def test_f5_libelle_interdit_non_marque_leve(lot):  # noqa: F811
    from controldone.guardrails import FormulationInterdite
    from controldone.rapport.export import findings_lot_json

    rd = _rd(lot)
    c0 = rd.findings.constats[0].model_copy(update={"libelle": "une fraude fictive"})
    rd.findings = rd.findings.model_copy(update={"constats": [c0]})
    with pytest.raises(FormulationInterdite):
        findings_lot_json([rd])


def test_f14_tableur_montants_sans_float(lot, tmp_path):  # noqa: F811
    from decimal import Decimal

    from openpyxl import load_workbook

    from controldone.rapport.export import ecrire_xlsx
    from controldone.rapport.vue import construire_vue

    rd = _rd(lot)
    rd.resultats = [r for r in rd.resultats if r.constat is None]
    rd.resultats.append(_resultat(rd, "C1", "ecart_certain", "1234.57", "recouvrable"))
    v = construire_vue([rd], rd.profil)
    p = ecrire_xlsx(v, [rd], tmp_path / "x.xlsx")
    ws = load_workbook(p)["Dossiers"]
    entetes = [c.value for c in ws[1]]
    val = ws.cell(row=2, column=entetes.index("Recouvrable certain (EUR)") + 1).value
    assert Decimal(str(val)) == Decimal("1234.57")
    # Règle « Decimal partout » (AGENTS_RULES 8) : aucune conversion explicite en float dans l'export.
    import inspect

    import controldone.rapport.export as export

    assert "float(" not in inspect.getsource(export)


# --- F12 : rejeu identique octet pour octet (hors horodatages) -----------------------------------------

def test_f12_rejeu_identique_et_annee_du_lot(lot, monkeypatch):  # noqa: F811
    from datetime import UTC, datetime

    import controldone.pipeline as pl
    from controldone.findings_io import findings_json_rejeu

    a = _rd(lot)
    b = _rd(lot)
    assert a.findings.execution.duree_s is not None  # Annexe C : le champ reste dans findings.json
    assert findings_json_rejeu(a.findings) == findings_json_rejeu(b.findings)
    assert "duree_s" not in findings_json_rejeu(a.findings)
    # La référence D-AAAA-NNNNN suit la date du lot, pas l'horloge au moment du traitement.
    monkeypatch.setattr(pl, "horodatage", lambda: datetime(2031, 1, 1, tzinfo=UTC))
    r = traiter_lot(lot / "docs", PROFIL, [], options=OptionsPipeline(seed=1), composants=_composants())[0]
    assert not r.dossier.reference.startswith("D-2031-")

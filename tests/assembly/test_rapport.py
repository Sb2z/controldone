"""Rapport de diagnostic (§18) : rendu HTML/PDF/JSON/XLSX et garde-fous (§3) — jeu de démonstration fictif."""

from __future__ import annotations

import json
import re
from decimal import Decimal

import pytest

from controldone.guardrails import AVERTISSEMENT, PHRASE_RENVOI, check_text
from controldone.rapport.html import texte_visible


@pytest.fixture(scope="module")
def demo(tmp_path_factory):
    from controldone.demo import executer_demo

    base = tmp_path_factory.mktemp("demo")
    # moteur « demo » : doubles de démonstration, résultat indépendant des autres équipes
    return executer_demo(base / "out", racine=base / "jeu", moteur="demo")


def _normaliser(t: str) -> str:
    return re.sub(r"\s+", " ", t.replace(" ", " ")).strip()


def test_constats_attendus_du_jeu_fictif(demo):
    v = demo.vue
    par_dossier = {d.reference: {(c.controle_id, c.niveau_code) for c in d.constats} for d in v.dossiers}
    assert ("C3", "ecart_certain") in par_dossier["D-2026-00001"]
    assert ("D3", "ecart_certain") in par_dossier["D-2026-00001"]
    assert ("B1", "ecart_certain") in par_dossier["D-2026-00002"]
    assert ("A12", "a_verifier") in par_dossier["D-2026-00002"]
    assert par_dossier["D-2026-00003"] == set()
    assert [d.statut_code for d in v.dossiers] == ["ecart_certain", "ecart_certain", "conforme"]


def test_totaux_par_nature_jamais_additionnes(demo):
    v = demo.vue
    recouvrables = [c.montant_valeur for c in v.constats() if c.nature_code == "recouvrable"
                    and c.niveau_code == "ecart_certain"]
    assert sum(recouvrables, Decimal(0)) == Decimal("2386.28")
    assert v.recouvrable_certain.replace(" ", " ") == "2 386,28 EUR"
    assert v.ecarts_calcul_nb == 1 and v.ecarts_calcul.replace(" ", " ") == "20,00 EUR"
    assert v.nb_renvois == 1
    # la note de renvoi n'a jamais de montant
    for c in v.renvois:
        assert c.montant_valeur is None and c.montant == "—"
        assert PHRASE_RENVOI in c.prochaine_action or PHRASE_RENVOI in c.libelle


def test_html_sobre_sans_ressource_externe(demo):
    html = demo.html.read_text(encoding="utf-8")
    assert "<script" not in html.lower()
    assert not re.search(r"""(?:src|href)\s*=\s*["']https?://""", html)
    assert "@media print" in html and "@page" in html
    texte = texte_visible(html)
    assert check_text(texte) == []
    assert _normaliser(AVERTISSEMENT) in _normaliser(texte)
    assert "DONNÉES FICTIVES" in texte
    assert _normaliser(PHRASE_RENVOI) in _normaliser(texte)
    # preuves côte à côte avec rognage de page
    assert html.count("data:image/png;base64,") >= 4
    for titre in ("Synthèse", "Prochaines actions", "Constats par nature", "Tableau des dossiers", "Fiches dossiers",
                  "Points à faire vérifier par un professionnel", "Documents non lus ou non reconnus",
                  "Méthode et tolérances", "Avertissement"):
        assert titre in texte


def test_pdf_avertissement_et_bandeau_sur_chaque_page(demo):
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(demo.pdf))
    try:
        assert len(pdf) >= 5
        textes = [_normaliser(pdf[i].get_textpage().get_text_range()) for i in range(len(pdf))]
    finally:
        pdf.close()
    for i, t in enumerate(textes, start=1):
        assert "Ce document est un contrôle technique de cohérence entre documents" in t, i
        assert "ne préjugent pas des sommes légalement dues" in t, i
        assert "DONNÉES FICTIVES" in t, i
        assert re.search(rf"Page {i} / {len(textes)}", t), i
        assert check_text(t) == [], i
    tout = " ".join(textes)
    assert "Rapport de diagnostic" in textes[0] and "ATELIERS DÉMO FICTIF SAS" in textes[0]
    assert "2 386,28 EUR" in tout


def test_tous_les_textes_de_la_vue_passent_les_garde_fous(demo):
    v = demo.vue
    textes = [v.titre, v.mention_validation, *v.limites, *(a.titre for a in v.actions),
              *(x for a in v.actions for x in a.details), *(n.libelle + n.note for n in v.table_nature)]
    for c in [*v.constats(), *v.renvois]:
        textes += [c.libelle, c.prochaine_action, *c.raisons, c.nature, c.controle_libelle]
    for t in textes:
        assert check_text(t) == [], t


def test_exports_json_et_xlsx(demo):
    data = json.loads(demo.json.read_text(encoding="utf-8"))
    assert data["schema"].startswith("controldone.rapport/1.")
    assert data["donnees_fictives"] is True and data["avertissement"] == AVERTISSEMENT
    assert len(data["dossiers"]) == 3
    findings = json.loads(demo.findings.read_text(encoding="utf-8"))
    assert isinstance(findings, list) and all(f["schema"] == "controldone.findings/1.0.0" for f in findings)
    from openpyxl import load_workbook

    wb = load_workbook(demo.xlsx)
    assert wb.sheetnames == ["Synthèse", "Dossiers", "Constats", "Contrôles", "Documents", "Non lus", "Méthode"]
    entetes = [c.value for c in wb["Constats"][1]]
    assert "Décision (valider / rejeter / à vérifier)" in entetes and "Montant en jeu (EUR)" in entetes
    assert wb["Constats"].max_row == 1 + 4


def test_libelle_bloque_remplace_par_mention_neutre(demo):
    from controldone.model import Constat, NatureMontant, Niveau, RaisonCode
    from controldone.rapport.vue import _constat_vue

    rd_vue = demo.vue
    assert rd_vue.constats_bloques == 0
    c = Constat(controle_id="B1", niveau=Niveau.a_verifier, raisons=[RaisonCode.ecart_sous_seuil],
                libelle="texte avec une fraude", nature_montant=NatureMontant.arithmetique_declaration,
                motif_blocage="formulation_interdite")

    class R:
        tolerance_appliquee = None
        seuil_certitude_applique = None

    class RD:
        documents, fichiers, chemins = {}, {}, {}

        class dossier:
            reference, id = "D-2026-99999", "dos_x"

    vue = _constat_vue(c, R(), RD())
    assert vue.bloque and check_text(vue.libelle) == []


def test_formulation_interdite_bloque_la_generation(demo, tmp_path, monkeypatch):
    from controldone.guardrails import FormulationInterdite
    from controldone.rapport import vue as module_vue

    monkeypatch.setattr(module_vue, "LIMITES_METHODE", ["Le taux erroné est corrigé."])
    from controldone.rapport import construire_vue, rendre_html, verifier_textes

    v = construire_vue([], _profil())
    with pytest.raises(FormulationInterdite):
        verifier_textes(v, rendre_html(v))


def _profil():
    from controldone.demo.donnees import profil_demo
    from controldone.referentiel_io import charger_profil_client

    return charger_profil_client(profil_demo())


def test_rapport_vide_ne_plante_pas(tmp_path):
    from controldone.rapport import generer_rapport

    s = generer_rapport([], _profil(), tmp_path / "vide")
    assert s.pdf.exists() and s.vue.nb_dossiers == 0
    assert "Aucune action" in texte_visible(s.html.read_text(encoding="utf-8"))

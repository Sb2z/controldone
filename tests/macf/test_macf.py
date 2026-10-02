"""Préparation MACF (CBAM) : sélection par code imprimé, agrégats, seuil arithmétique, brouillons, exports.

Toutes les données sont FICTIVES.
"""

from __future__ import annotations

import csv
from datetime import date
from decimal import Decimal

import pytest
from cryptography.fernet import Fernet

from controldone.auth.roles import Acteur, Role
from controldone.guardrails import AVERTISSEMENT, PHRASE_RENVOI, check_text
from controldone.macf import (
    DONNEES_A_DEMANDER,
    MENTION_PREPARATION,
    StatutCode,
    agreger,
    brouillons_demandes,
    charger_liste,
    ecrire_pack,
    lignes_depuis_scope,
    masse_kg,
    preparer_pack,
    proposer_brouillons,
    selectionner_lignes,
    synthese_seuil,
)
from controldone.model import ArticleDeclaration, Partie
from controldone.testing import declaration, dossier_pour, facture_commerciale, vs

LISTE = charger_liste()


def _article(doc_id: str, n: int, code: str, masse: str | None, origine: str = "CN", unite: str = "KGM"):
    base = f"declaration.articles[{n - 1}]"
    return ArticleDeclaration(
        numero_article=vs(f"{base}.numero_article", str(n), document_id=doc_id),
        code_marchandise=vs(f"{base}.code_marchandise", code, document_id=doc_id, page=2),
        description=vs(f"{base}.description", f"MARCHANDISE FICTIVE {n}", document_id=doc_id),
        pays_origine=vs(f"{base}.pays_origine", origine, document_id=doc_id),
        masse_nette=(vs(f"{base}.masse_nette", masse, document_id=doc_id, unite=unite) if masse else None),
    )


def _dossier(num: int, date_acc: str, articles: list[tuple[str, str | None, str]], vendeur: str):
    doc_id = f"doc_dec_{num}"
    dec = declaration(id=doc_id, mrn=f"26FRFIC00000000{num:03d}",
                      date_acceptation=vs("declaration.date_acceptation", date_acc, document_id=doc_id),
                      articles=[_article(doc_id, i + 1, c, m, o) for i, (c, m, o) in enumerate(articles)])
    fc = facture_commerciale(id=f"doc_fc_{num}", vendeur=Partie(
        nom=vs("facture_commerciale.vendeur.nom", vendeur, document_id=f"doc_fc_{num}")))
    dos = dossier_pour([dec, fc], id=f"dos_{num}")
    dos.reference = f"D-2026-{num:05d}"
    return dos, [dec, fc]


@pytest.fixture
def lignes():
    out = []
    for dos, docs in (
        _dossier(1, "2026-02-10", [("7208 51 98", "12000.000", "TR"), ("8471300000", "50.000", "CN")],
                 "ACIERIE FICTIVE ANADOLU"),
        _dossier(2, "2026-05-03", [("7601 10 00", "8500.500", "IN"), ("7202", "100.000", "CN")],
                 "ALU FICTIF LTD"),
        _dossier(3, "2026-08-21", [("2804 10 00", "900.000", "NO"), ("7318 15 90", None, "CN")],
                 "VISSERIE FICTIVE CO"),
        _dossier(4, "2025-12-30", [("7308 90 98", "3000.000", "CN")], "ACIERIE FICTIVE ANADOLU"),
    ):
        out += selectionner_lignes(dos, docs, liste=LISTE)
    return out


# --- liste des codes ------------------------------------------------------------------------------------

def test_liste_versionnee_sourcee_et_a_verifier():
    assert LISTE.version and LISTE.source_url.startswith("https://eur-lex.europa.eu/")
    assert LISTE.consulte_le == "2026-10-02" and LISTE.a_verifier
    assert LISTE.seuil_tonnes == "50"
    assert {s for s in LISTE.secteurs} == {"ciment", "electricite", "engrais", "fer_acier", "aluminium", "hydrogene"}
    assert LISTE.secteurs["electricite"].hors_cumul_50t and LISTE.secteurs["hydrogene"].hors_cumul_50t


@pytest.mark.parametrize(("code", "statut", "secteur"), [
    ("7208 51 98", StatutCode.dans_liste, "fer_acier"),
    ("72.08.51.98.00", StatutCode.dans_liste, "fer_acier"),
    ("7204 49 00", StatutCode.hors_liste, None),          # déchets : exclusion de l'entrée « 72 »
    ("7202 21 00", StatutCode.hors_liste, None),          # ferro-silicium : exclusion « 7202 2 »
    ("7202 11 20", StatutCode.dans_liste, "fer_acier"),   # ferro-manganèse : non exclu
    ("7202", StatutCode.a_preciser, "fer_acier"),         # trop court face aux exclusions
    ("2523", StatutCode.a_preciser, "ciment"),
    ("2523 29 00", StatutCode.dans_liste, "ciment"),
    ("3105 60 00", StatutCode.hors_liste, None),
    ("3105 20 10", StatutCode.dans_liste, "engrais"),
    ("2716 00 00", StatutCode.dans_liste, "electricite"),
    ("7616 99 90", StatutCode.dans_liste, "aluminium"),
    ("7615 10 10", StatutCode.hors_liste, None),
    ("8471 30 00", StatutCode.hors_liste, None),
    ("", StatutCode.illisible, None),
    (None, StatutCode.illisible, None),
])
def test_classement_par_code_imprime(code, statut, secteur):
    c = LISTE.classer(code)
    assert c.statut is statut
    assert (c.secteur.id if c.secteur else None) == (secteur if statut is not StatutCode.hors_liste else None)


def test_masse_kg():
    assert masse_kg(vs("declaration.articles[0].masse_nette", "12.5", unite="TNE")) == Decimal("12500.000")
    assert masse_kg(vs("declaration.articles[0].masse_nette", "1250.5", unite="KGM")) == Decimal("1250.500")
    assert masse_kg(vs("declaration.articles[0].masse_nette", "10", unite="LTR")) is None
    assert masse_kg(None) is None


# --- sélection ----------------------------------------------------------------------------------------------

def test_selection_par_code_imprime_et_provenance(lignes):
    codes = [li.code_imprime for li in lignes]
    assert "8471300000" not in codes  # hors liste
    assert codes == ["7208 51 98", "7601 10 00", "7202", "2804 10 00", "7318 15 90", "7308 90 98"]
    li = lignes[0]
    assert li.fournisseur == "ACIERIE FICTIVE ANADOLU" and li.pays_origine == "TR"
    assert li.mrn == "26FRFIC00000000001" and li.page == 2 and li.periode == "2026-T1"
    assert li.verification == "à faire vérifier" and li.installation is None
    assert next(x for x in lignes if x.code_imprime == "7202").statut_code is StatutCode.a_preciser
    assert next(x for x in lignes if x.code_imprime == "2804 10 00").hors_cumul_50t


def test_agregats_par_code_origine_fournisseur_periode(lignes):
    ags = agreger(lignes, annee=2026)
    assert {a.periode for a in ags} == {"2026-T1", "2026-T2", "2026-T3"}
    alu = next(a for a in ags if a.code_imprime == "7601 10 00")
    assert alu.masse_nette_kg == Decimal("8500.500") and alu.pays_origine == "IN" and alu.fournisseur == "ALU FICTIF LTD"
    assert alu.installation == "à demander au fournisseur"
    sans = next(a for a in ags if a.code_imprime == "7318 15 90")
    assert sans.lignes_sans_masse == 1 and sans.masse_nette_kg == 0
    assert all(isinstance(a.masse_nette_kg, Decimal) for a in ags)


# --- seuil : arithmétique seulement --------------------------------------------------------------------

def test_seuil_arithmetique_avec_renvoi(lignes):
    s = synthese_seuil(lignes, 2026)
    # 12 000 + 8 500,5 + 100 (code à préciser) ; hydrogène exclu ; 7318 sans masse non compté ; 2025 hors année
    assert s.masse_cumulee_kg == Decimal("20600.500")
    assert s.masse_cumulee_t == Decimal("20.601") and s.difference_t == Decimal("29.399")
    assert s.lignes_hors_cumul == 1 and s.lignes_sans_masse == 1
    assert PHRASE_RENVOI in s.textes
    assert "29,399 t" in s.texte
    assert check_text(s.texte) == []
    for mot in ("exempt", "soumis", "assujetti", "obligé", "redevable"):
        assert mot not in s.texte.lower()


def test_seuil_depasse_reste_un_calcul():
    dos, docs = _dossier(9, "2026-03-01", [("7601 20 20", "60000.000", "AE")], "ALU FICTIF LTD")
    s = synthese_seuil(selectionner_lignes(dos, docs, liste=LISTE), 2026)
    assert s.difference_t == Decimal("-10.000")
    assert "60,000 t - 50,000 t = 10,000 t" in s.texte.replace(" ", " ")
    assert PHRASE_RENVOI in s.textes and check_text(s.texte) == []


# --- demandes aux fournisseurs ------------------------------------------------------------------------

def test_liste_des_donnees_a_demander():
    ids = {d.id for d in DONNEES_A_DEMANDER}
    assert {"installation", "emissions_directes", "prix_carbone"} <= ids
    for d in DONNEES_A_DEMANDER:
        assert check_text(f"{d.libelle} {d.precision} {d.libelle_en}") == []


def test_brouillons_un_par_fournisseur_jamais_envoyes(lignes):
    bs = brouillons_demandes(lignes, annee=2026, client="CLIENT FICTIF SAS")
    assert [b["fournisseur"] for b in bs] == ["ACIERIE FICTIVE ANADOLU", "ALU FICTIF LTD", "VISSERIE FICTIVE CO"]
    for b in bs:
        assert b["destinataires"] == [] and b["destinataire_role"] == "client"
        assert PHRASE_RENVOI in b["corps"] and AVERTISSEMENT in b["corps"] and MENTION_PREPARATION in b["corps"]
        assert "émissions" in b["corps"].lower() and "installation" in b["corps"].lower()
        assert check_text(b["objet"]) == [] and check_text(b["corps"]) == []
    # 2025 : la ligne de décembre 2025 n'entre pas dans les brouillons 2026
    assert "7308 90 98" not in bs[0]["corps"]


@pytest.fixture
def db(tmp_path, monkeypatch):
    from controldone.storage import Database

    monkeypatch.setenv("CONTROLDONE_ENV", "test")
    monkeypatch.setenv("CONTROLDONE_MASTER_KEY", Fernet.generate_key().decode())
    monkeypatch.setenv("CONTROLDONE_DATA_DIR", str(tmp_path / "var"))
    base = Database(f"sqlite:///{tmp_path}/macf.db")
    base.creer_schema()
    with base.operateur(Acteur("usr_fondateur", Role.fondateur)) as op:
        op.creer_client("cli_macf", "CLIENT MACF FICTIF")
    yield base
    base.fermer()


def test_brouillons_deposes_dans_la_file_sans_envoi(db, lignes):
    from controldone.outbox import FileSortante, StatutAction

    file = FileSortante(db)
    bs = brouillons_demandes(lignes, annee=2026, client="CLIENT FICTIF SAS")
    actions = proposer_brouillons(file, bs, Acteur.systeme("tests"), tenant_id="cli_macf")
    assert len(actions) == 3
    assert all(a.statut is StatutAction.brouillon and a.motif_blocage is None for a in actions)
    assert all(a.envoye_le is None for a in actions)
    # idempotent
    again = proposer_brouillons(file, bs, Acteur.systeme("tests"), tenant_id="cli_macf")
    assert [a.id for a in again] == [a.id for a in actions]


def test_lignes_depuis_les_dossiers_traites_du_client(db):
    systeme = Acteur.systeme("tests")
    with db.tenant("cli_macf", systeme) as sc:
        for n, d, arts, v in ((1, "2026-02-10", [("7208 51 98", "12000.000", "TR")], "ACIERIE FICTIVE"),
                              (2, "2026-04-02", [("6403 99 93", "40.000", "VN")], "CHAUSSURES FICTIVES")):
            dos, docs = _dossier(n, d, arts, v)
            sc.enregistrer_dossier(dos, documents=docs)
    with db.tenant("cli_macf", systeme, lecture=True) as sc:
        ls = lignes_depuis_scope(sc, liste=LISTE)
    assert [li.code_imprime for li in ls] == ["7208 51 98"]
    assert ls[0].masse_nette_kg == Decimal("12000.000")


# --- exports -----------------------------------------------------------------------------------------------

def test_pack_csv_xlsx_pdf(tmp_path, lignes):
    pack = preparer_pack(lignes, annee=2026, client="CLIENT FICTIF SAS", date_preparation=date(2026, 10, 2))
    assert len(pack.lignes) == 5 and all(li.annee == 2026 for li in pack.lignes)
    for t in pack.textes:
        assert check_text(t) == []
    sorties = ecrire_pack(pack, tmp_path / "pack")
    csv_lignes = sorties["csv"][0].read_text(encoding="utf-8-sig")
    assert MENTION_PREPARATION.upper() in csv_lignes and PHRASE_RENVOI in csv_lignes
    rangs = list(csv.reader(csv_lignes.splitlines(), delimiter=";"))
    assert sum(1 for r in rangs if r and r[-1] == "à faire vérifier") == 5

    from openpyxl import load_workbook

    wb = load_workbook(sorties["xlsx"])
    assert wb.sheetnames == ["Synthèse", "Lignes", "Agrégats"]
    assert wb["Synthèse"]["A1"].value == MENTION_PREPARATION.upper()
    assert wb["Lignes"]["A1"].value == MENTION_PREPARATION.upper()

    from pypdf import PdfReader

    texte = " ".join(p.extract_text() for p in PdfReader(str(sorties["pdf"])).pages)
    texte = " ".join(texte.split())
    assert MENTION_PREPARATION.upper() in texte
    assert "Ce point relève d'une appréciation réglementaire" in texte
    assert check_text(texte) == []


def test_injection_de_formule_neutralisee(tmp_path):
    dos, docs = _dossier(7, "2026-01-15", [("7208 51 98", "10.000", "CN")], "=HYPERLINK(\"http://x\")")
    pack = preparer_pack(selectionner_lignes(dos, docs, liste=LISTE), annee=2026, client="CLIENT FICTIF",
                         date_preparation=date(2026, 10, 2))
    sorties = ecrire_pack(pack, tmp_path)
    assert "'=HYPERLINK" in sorties["csv"][0].read_text(encoding="utf-8-sig")
    from openpyxl import load_workbook

    ws = load_workbook(sorties["xlsx"])["Lignes"]
    assert all(c.data_type != "f" for r in ws.iter_rows() for c in r)

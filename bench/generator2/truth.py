"""Construction de truth.json (SPEC §19.3, Annexe B étendue) et auto-contrôle « chaque valeur de
vérité est imprimée » (re-rendu natif puis recherche dans le texte)."""

from __future__ import annotations

import io
import json
import re
from decimal import Decimal as D
from pathlib import Path

from .build import ANNEXE_A, COMPOSANTE, NATURE, STRUCTURED_FORMATS, expected_level
from .model import FAMCAP, LAYCAP
from .render_ft import _multi_mrn, tr_ref
from .ext import VAT_INCLUSIVE_FAMILIES
from .util import cur_decimals, output_version, date_candidates, fmt_hs, fmt_num, q2, s2

PAY_NORM = {"A": "comptant", "E": "differe", "G": "autoliquide"}
SCHEMA_PATH = Path(__file__).with_name("truth.schema.json")


def _m(v, devise="EUR"):
    if v is None:
        return None
    dec = cur_decimals(devise)
    return str(D(v).quantize(D(1).scaleb(-dec))) if dec else str(int(D(v).to_integral_value()))


def _d3(v):
    return None if v is None else str(D(v).quantize(D("0.001")))


def _qty(v):
    if v is None:
        return None
    v = D(v)
    return str(int(v)) if v == v.to_integral_value() else format(v.normalize(), "f")


def _rate(v):
    if v is None:
        return None
    s = format(D(v).normalize(), "f")
    return s


def _lang_doc(doc):
    k = doc["kind"]
    if k in ("ci",):
        return doc["lang"]
    if k == "dec":
        return "en" if doc["layout"] in ("M3", "M9") else "fr"
    if k in ("ft", "av"):
        return "fr_en" if doc["lang"] == "fr_en" else doc["lang"]
    return doc.get("lang", "fr") if doc.get("lang") != "fr_en" else "fr_en"


# ---------------------------------------------------------------------------
# Valeurs de vérité par document
# ---------------------------------------------------------------------------

def tv_ci(ci):
    dev = ci["devise"]
    lines = []
    for ln in ci["lines"]:
        hs = None
        if ci["hs_digits"]:
            style = ci["hs_style"] if ci.get("final_format") not in ("ubl",) else "plain"
            hs = fmt_hs(ln["hs10"], ci["hs_digits"], style)
        lines.append({"numero_ligne": ln["no"], "reference_article": ln["ref"], "description": ln["desc"],
                      "code_marchandise_imprime": hs, "quantite": _qty(ln["qty"]), "unite": ln["unit"],
                      "prix_unitaire": _rate(ln["pu"]), "montant_ligne": _m(ln["amount"], dev), "pays_origine": ln["origin"],
                      "masse_nette": None, "masse_brute": None})
    sub = {k: _m(v, dev) for k, v in ci["sub"].items()} if len(ci["sub"]) > 1 else {}
    return {"numero": ci["numero"], "date": ci["date"].isoformat(), "devise": dev, "total_facture": _m(ci["total"], dev),
            "total_imprime": True, "vendeur.nom": ci["supplier"]["nom"],
            "acheteur.nom": ci["acheteur"].get("nom_affiche", ci["acheteur"]["nom"]),
            "acheteur.tva": ci["acheteur"]["tva"], "incoterm": ci["incoterm"], "incoterm_lieu": ci["incoterm_lieu"],
            "ref_transport": ci["ref_transport"], "masse_nette_totale": _d3(ci["net_total"]),
            "masse_brute_totale": _d3(ci["gross_total"]), "nombre_colis": ci["colis"], "sous_totaux": sub, "lignes": lines}


def tv_dec(dec):
    dev = dec["devise"]
    lay = LAYCAP[dec["layout"]]
    arts = []
    for a in dec["articles"]:
        arts.append({"numero_article": a["no"], "code_marchandise": a["hs10"], "pays_origine": a["origin"],
                     "masse_nette": _d3(a["net"]), "masse_brute": _d3(a["gross"]), "montant_facture_article": _m(a["amount"], dev),
                     "valeur_statistique": s2(a["stat_value"]), "nombre_colis": a["colis"] if lay["colis_art"] else None,
                     "quantite_unite_supplementaire": _qty(a["qty"]) if (lay["qty"] and a["qty"] is not None) else None})
    taxes = []
    for t in dec["taxes"]:
        taxes.append({"article": t["article"], "type_taxe": t["type"], "categorie": t["cat"],
                      "base_montant": s2(t["base"]) if t["base"] is not None else None,
                      "base_quantite": _qty(t["base_qty"]) if t["base_qty"] is not None else None,
                      "taux": _rate(t["rate"]), "taux_nature": t["nature"], "montant": s2(t["montant"]),
                      "montant_a_payer": s2(t["a_payer"]), "paiement_normalise": PAY_NORM[t["mode"]]})
    ind = [r for r in dec["refs"] if r["code"] in ("1008", "FR7")]
    return {"mrn": dec["mrn"], "lrn": dec["lrn"], "version": dec["version"], "date_acceptation": dec["date"].isoformat(),
            "importateur.nom": dec["importateur"]["nom"], "importateur.tva": dec["importateur"]["tva"],
            "declarant.nom": dec["declarant"]["nom"], "devise_facture": dev, "montant_total_facture": _m(dec["montant"], dev),
            "taux_change": str(dec["rate"]) if dec["rate"] is not None else None,
            "taux_change_sens": dec["sens"] if dec["rate"] is not None else None,
            "incoterm": dec["incoterm"], "masse_brute_totale": _d3(dec["gross_total"]), "nombre_colis_total": dec["colis_total"],
            "nombre_articles": dec["n_articles"],
            "documents_references": [{"type_code": r["code"], "reference": r["ref"]} for r in dec["refs"]],
            "indices_autoliquidation": bool(dec["autoliq"]) and bool(ind),
            "articles": arts, "taxations": taxes,
            "totaux_par_type": {k: s2(v) for k, v in dec["type_totals"].items()} if lay["type_totals"] else None,
            "total_droits_taxes": s2(dec["total_droits_taxes"]), "total_a_payer": s2(dec["total_a_payer"])}


def _line_mrn_printed(doc, ln):
    fam = doc["family"]
    if not ln.get("mrn"):
        return False
    if fam in ("G1", "G2", "G7", "G8", "G13", "G15", "G16"):
        return True
    if fam == "G14":
        return ln["nature"].startswith("debours")
    if fam in ("G3", "G4", "G9", "G10"):
        return _multi_mrn(doc)
    if fam == "G5":
        return ln["nature"].startswith("debours")
    if fam == "G12":
        return all(x["nature"].startswith("debours") for x in doc["lines"])
    return False


def _line_vat_printed(doc, ln):
    fam = doc["family"]
    if fam in ("G2", "G3", "G10", "G14", "G15"):
        return False
    if fam == "G5":
        return not ln["nature"].startswith("debours")
    if fam == "G12":
        return not all(x["nature"].startswith("debours") for x in doc["lines"])
    return True


def _total_debours_printed(doc):
    if doc["kind"] == "av":
        return False
    fam = doc["family"]
    if not FAMCAP[fam]["total_debours"]:
        return False
    return any(ln["nature"].startswith("debours") for ln in doc["lines"])


def _net_printed(doc):
    return doc["kind"] != "av" and doc["family"] in ("G1", "G2", "G6", "G7", "G8", "G13", "G14", "G16")


def tv_ft(doc):
    lines = []
    for ln in doc["lines"]:
        vp = _line_vat_printed(doc, ln)
        pu_printed = doc["family"] not in ("G2", "G10", "G15")      # G2, G10, G15 : quantité et montant net seulement
        ttc_line = doc["family"] in VAT_INCLUSIVE_FAMILIES and ln["vat_rate"] > 0   # G13 : ligne TTC
        if ttc_line:
            pu_printed = False
        lines.append({"nature": ln["nature"], "libelle": ln["libelle"], "quantite": _qty(ln["qty"]),
                      "prix_unitaire": _rate(ln["pu"]) if pu_printed else None,
                      "montant_ht": s2(abs(ln["ht"])) if doc["kind"] == "av" else s2(ln["ht"]),
                      "taux_tva": _rate(ln["vat_rate"]) if vp else None, "montant_tva": s2(ln["vat"]) if (vp and doc["family"] not in ("G7", "G8")) else None,
                      "mrn": ln["mrn"] if _line_mrn_printed(doc, ln) else None,
                      "date_debut": ln["date_debut"].isoformat() if ln.get("date_debut") else None,
                      "date_fin": ln["date_fin"].isoformat() if ln.get("date_fin") else None})
        if doc["family"] in VAT_INCLUSIVE_FAMILIES:     # 2.1 (G13) : TVA par ligne non imprimée ; lignes taxables TTC
            lines[-1]["montant_tva"] = None
            if ttc_line:
                lines[-1]["montant_ht"] = None
                lines[-1]["montant_ttc"] = s2(abs(ln["ht"]) + abs(ln["vat"]))
    if doc["kind"] == "av":
        return {"numero": doc["numero"], "date": doc["date"].isoformat(), "emetteur.tva": doc["emetteur"]["tva"],
                "refs_facture_origine": list(doc["refs_facture_origine"]), "refs_mrn": list(doc["refs_mrn"]),
                "refs_transport": [tr_ref(r, doc.get("transport_ref_style", "raw")) for r in doc["refs_transport"]],
                "lignes": [{"nature": x["nature"], "libelle": x["libelle"], "montant_ht": x["montant_ht"],
                            **({"montant_ttc": x["montant_ttc"]} if "montant_ttc" in x else {})} for x in lines],
                "total_credite_ht": s2(doc["total_ht"]), "total_tva": s2(doc["total_tva"]), "total_credite_ttc": s2(doc["total_ttc"]),
                "devise": "EUR", "motif": doc.get("motif")}
    return {"numero": doc["numero"], "date": doc["date"].isoformat(), "emetteur.nom": doc["emetteur"]["nom"],
            "emetteur.tva": doc["emetteur"]["tva"], "client_facture.nom": doc["client"]["nom"], "client_facture.tva": doc["client"]["tva"],
            "refs_mrn": list(doc["refs_mrn"]),
            "refs_transport": [tr_ref(r, doc.get("transport_ref_style", "raw")) for r in doc["refs_transport"]],
            "devise": "EUR", "lignes": lines,
            "total_debours": s2(doc["total_debours"]) if _total_debours_printed(doc) else None,
            "total_ht": s2(doc["total_ht"]), "total_tva": s2(doc["total_tva"]), "total_ttc": s2(doc["total_ttc"]),
            "net_a_payer": s2(doc["net"]) if _net_printed(doc) else None, "est_releve": bool(doc.get("est_releve"))}


def tv_sup(doc):
    ci = doc.get("ci")
    if doc["sous_type"] == "liste_colisage":
        return {"ref_transport": ci["ref_transport"], "nombre_colis": ci["colis"], "masse_brute": _d3(ci["gross_total"]),
                "refs_facture": [ci["numero"]]}
    if doc["sous_type"] == "titre_transport":
        return {"ref_transport": ci["transport"]["ref"], "nombre_colis": ci["colis"], "masse_brute": _d3(ci["gross_total"])}
    if doc["sous_type"] == "pre_alerte":
        return {"ref_transport": ci["ref_transport"]}
    return {}


def truth_values(doc):
    k = doc["kind"]
    if k == "ci":
        return tv_ci(doc)
    if k == "dec":
        return tv_dec(doc)
    if k in ("ft", "av"):
        return tv_ft(doc)
    return tv_sup(doc)


# ---------------------------------------------------------------------------
# truth.json
# ---------------------------------------------------------------------------

def build_truth(dos, files, seed) -> dict:
    sp = dos.spec
    present = [d for d in dos.order if d not in dos.removed]
    documents = []
    tvs = {}
    for d in present:
        doc = dos.docs[d]
        documents.append({"doc_id": d, "type": doc["type"], "sous_type": doc.get("sous_type"), "format": doc["final_format"],
                          "degradation": doc["deg"], "degradation_mode": doc["mode"], "file": doc["file"], "pages": doc["pages"],
                          "transitaire_template": doc.get("family") if doc["kind"] in ("ft", "av") else None,
                          "declaration_layout": doc["layout"] if doc["kind"] == "dec" else None,
                          "ci_layout": doc["layout"] if doc["kind"] == "ci" else None,
                          "language": _lang_doc(doc)})
        tvs[d] = truth_values(doc)
    errors = []
    tot_c = D(0)
    tot_v = D(0)
    for i, e in enumerate(dos.errors, 1):
        ctrl = e["ctrl"]
        lvl = expected_level(dos, e)
        nat = NATURE[ctrl]
        amount = None
        if nat not in ("aucun", "renvoi") and e["amount"] is not None:
            amount = s2(e["amount"])
        ent = {"error_id": f"{dos.did}-E{i}", "control_id": ctrl, "accepted_control_ids": list(ANNEXE_A[ctrl]),
               "expected_level": lvl, "expected_amount_eur": amount, "amount_nature": nat,
               "composante": COMPOSANTE.get(ctrl), "documents": [d for d in e["docs"] if d in present] or e["docs"][:1],
               "fields": e["fields"], "injection": e["code"], "description": e["desc"]}
        if ctrl in ("F2", "F3", "F4", "F5"):
            ent["other_dossiers"] = list(e["li"].get("other", []))
        errors.append(ent)
        if nat == "recouvrable" and amount is not None and D(amount) > 0 and ctrl not in ("E6",):
            if lvl == "ecart_certain":
                tot_c += D(amount)
            else:
                tot_v += D(amount)
    traps = []
    for i, t in enumerate(dos.traps, 1):
        docs = [d for d in t["docs"] if d in present]
        if not docs:
            continue
        # un piège ne doit pas contredire une erreur injectée appariable (même contrôle accepté, mêmes documents)
        if any(t["ctrl"] in e["accepted_control_ids"] and set(docs) & set(e["documents"]) for e in errors):
            continue
        traps.append({"trap_id": f"{dos.did}-T{i}", "control_id": t["ctrl"], "max_level": t["max"], "documents": docs,
                      "description": t["desc"]})
    if any(e["control_id"] == "P1" for e in errors):
        outcome = "document_manquant"
    elif any(e["expected_level"] == "ecart_certain" for e in errors):
        outcome = "ecart_certain"
    elif errors:
        outcome = "a_verifier"
    else:
        outcome = "conforme"
    fts = [dos.docs[d] for d in present if dos.docs[d]["kind"] == "ft"]
    decs = [dos.docs[d] for d in present if dos.docs[d]["kind"] == "dec"]
    tags = list(dict.fromkeys(["avoirs" if x == "avoir" else x for x in dos.tags] + scenario_tags(dos, present, files)))
    return {
        "schema": "controldone.bench.truth/1.0.0",
        "dossier_id": dos.did, "split": sp["split"], "client_id": sp["client_id"],
        "generator": "bench.generator2", "generator_version": output_version(bool(sp.get("ext"))), "seed": seed,
        "transitaire_template": sp["family"] if fts else None,
        "declaration_layout": decs[0]["layout"] if decs else None,
        "degradation": dos.deg_dossier,
        "scenario_tags": tags,
        "files": [{"path": f["path"], "sha256": f["sha256"], "pages": f["pages"]} for f in files],
        "documents": documents,
        "truth_values": tvs,
        "expected_links": [ln for ln in dos.links if ln["from"] in present and ln["to"] in present],
        "injected_errors": errors,
        "traps": traps,
        "expected_outcome": outcome,
        "expected_totals": {"recouvrable_certain_eur": s2(tot_c), "recouvrable_a_verifier_eur": s2(tot_v)},
    }


def scenario_tags(dos, present, files):
    t = []
    if any(len(f["docs"]) > 1 for f in files):
        t.append("pdf_fusionne")
    if any(dos.docs[d]["kind"] == "dec" and dos.docs[d]["autoliq"] for d in present):
        t.append("autoliquidation")
    if dos.attrs.get("devise", "EUR") != "EUR":
        t.append("devise_etrangere")
        if dos.attrs["devise"] in ("JPY", "KRW"):
            t.append("devise_sans_decimales")
    if any(dos.docs[d]["kind"] == "av" for d in present):
        t.append("avoirs")
    if dos.spec["kind"] == "h7":
        t.append("petits_envois")
    if any(dos.docs[d]["final_format"] in ("ubl", "cii", "factur_x") for d in present):
        t.append("facture_electronique")
    if any(dos.docs[d]["final_format"] == "xlsx" for d in present):
        t.append("tableur")
    if any(dos.docs[d]["kind"] == "ft" and len(dos.docs[d]["refs_mrn"]) > 1 for d in present):
        t.append("facture_transitaire_multi_mrn")
    if len([d for d in present if dos.docs[d]["kind"] == "dec" and not d.endswith("_copie")]) > 1:
        t.append("multi_declarations")
    if any(dos.docs[d].get("hidden_instr") for d in present):
        t.append("consigne_cachee")
    modes = {dos.docs[d]["mode"] for d in present}
    t += sorted(f"mode_{m}" for m in modes if m not in ("native", "structured"))
    return t


# ---------------------------------------------------------------------------
# Auto-contrôle : chaque valeur de vérité est imprimée
# ---------------------------------------------------------------------------

_WS = re.compile(r"[\s   ]+")


def norm_text(s: str) -> str:
    return _WS.sub("", s).lower()


def extract_text(data: bytes, fmt: str) -> str:
    if fmt == "pdf":
        import pypdf
        r = pypdf.PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in r.pages)
    if fmt == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True)
        out = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if v is None:
                        continue
                    if isinstance(v, float):
                        d = D(repr(v))
                        out.append(format(d, "f"))
                        out.append(format(d.quantize(D("0.01")), "f"))
                        out.append(format(d.quantize(D("0.001")), "f"))
                    else:
                        out.append(str(v))
        return " | ".join(out)
    if fmt == "csv_declaration":
        return data.decode("cp1252")
    if fmt == "eml":
        import email
        import email.policy
        msg = email.message_from_bytes(data, policy=email.policy.default)
        return str(msg["subject"]) + "\n" + msg.get_content()
    return data.decode("utf-8")


def _num_cands(v, decs):
    v = abs(D(v))
    out = set()
    for d in decs:
        for st in ("en", "de", "frs", "ch", "plain", "plainc"):
            out.add(fmt_num(v, st, d))
    return out


def _cands(kind, val):
    if kind == "amount2":
        v = D(val)
        # montant entier imprimé sans décimales (droits arrondis à l'euro, export brut) : même valeur
        return _num_cands(val, [2, 0] if v == v.to_integral_value() else [2])
    if kind == "amount0":
        return _num_cands(val, [0])
    if kind == "pu":
        v = D(val)
        s = format(v.normalize(), "f")
        dec = len(s.split(".")[1]) if "." in s else 0
        return _num_cands(v, sorted({dec, max(2, dec), 4, 2}))
    if kind == "qty":
        return _num_cands(val, [0, 1, 2, 3])
    if kind == "mass":
        return _num_cands(val, [3, 2, 1])
    if kind == "rate":
        return _num_cands(val, [0, 1, 2, 3, 4, 5])
    if kind == "date":
        from datetime import date as _date
        d = _date.fromisoformat(val)
        return set(date_candidates(d)) | {d.strftime("%Y%m%d"), d.strftime("%d/%m/%Y"), d.strftime("%d %b %Y")}
    if kind == "int":
        return {str(val)}
    return {str(val)}


def checks_for(doc, tv):
    """Liste (champ, type, valeur) de ce qui doit être imprimé."""
    k = doc["kind"]
    out = []

    def add(f, kind, v):
        if v is not None and v != "" and not (kind.startswith("amount") and D(v) == 0):
            out.append((f, kind, v))

    if k == "ci":
        dev = tv["devise"]
        ak = "amount0" if cur_decimals(dev) == 0 else "amount2"
        add("numero", "text", tv["numero"])
        add("date", "date", tv["date"])
        add("devise", "text", dev)
        add("total_facture", ak, tv["total_facture"])
        add("acheteur.tva", "text", tv["acheteur.tva"])
        add("incoterm", "text", tv["incoterm"])
        add("masse_brute_totale", "mass", tv["masse_brute_totale"])
        add("nombre_colis", "int", tv["nombre_colis"])
        add("ref_transport", "text", tv["ref_transport"])
        for i, ln in enumerate(tv["lignes"]):
            add(f"lignes[{i}].montant_ligne", ak, ln["montant_ligne"])
            add(f"lignes[{i}].quantite", "qty", ln["quantite"])
            add(f"lignes[{i}].pays_origine", "text", ln["pays_origine"])
            add(f"lignes[{i}].reference_article", "text", ln["reference_article"])
            if ln["code_marchandise_imprime"]:
                add(f"lignes[{i}].code_marchandise_imprime", "text", ln["code_marchandise_imprime"])
    elif k == "dec":
        dev = tv["devise_facture"]
        ak = "amount0" if cur_decimals(dev) == 0 and dev != "EUR" else "amount2"
        add("mrn", "text", tv["mrn"])
        add("date_acceptation", "date", tv["date_acceptation"])
        add("importateur.tva", "text", tv["importateur.tva"])
        add("devise_facture", "text", dev)
        add("montant_total_facture", ak, tv["montant_total_facture"])
        add("taux_change", "rate", tv["taux_change"])
        add("incoterm", "text", tv["incoterm"])
        add("nombre_articles", "int", tv["nombre_articles"])
        add("total_a_payer", "amount2", tv["total_a_payer"])
        add("total_droits_taxes", "amount2", tv["total_droits_taxes"])
        add("masse_brute_totale", "mass", tv["masse_brute_totale"])
        for r in tv["documents_references"]:
            add("documents_references", "text", r["reference"])
        for i, a in enumerate(tv["articles"]):
            add(f"articles[{i}].code_marchandise", "text", a["code_marchandise"])
            add(f"articles[{i}].pays_origine", "text", a["pays_origine"])
            add(f"articles[{i}].masse_nette", "mass", a["masse_nette"])
            add(f"articles[{i}].masse_brute", "mass", a["masse_brute"])
            add(f"articles[{i}].montant_facture_article", ak, a["montant_facture_article"])
            add(f"articles[{i}].quantite_unite_supplementaire", "qty", a["quantite_unite_supplementaire"])
        for i, t in enumerate(tv["taxations"]):
            add(f"taxations[{i}].type_taxe", "text", t["type_taxe"])
            add(f"taxations[{i}].base_montant", "amount2", t["base_montant"])
            add(f"taxations[{i}].base_quantite", "qty", t["base_quantite"])
            add(f"taxations[{i}].taux", "rate", t["taux"])
            add(f"taxations[{i}].montant", "amount2", t["montant"])
        for code, v in (tv.get("totaux_par_type") or {}).items():
            add(f"totaux_par_type.{code}", "amount2", v)
    elif k in ("ft", "av"):
        add("numero", "text", tv["numero"])
        add("date", "date", tv["date"])
        add("emetteur.tva", "text", tv["emetteur.tva"])
        if k == "ft":
            add("client_facture.tva", "text", tv["client_facture.tva"])
            add("total_debours", "amount2", tv["total_debours"])
            add("total_ht", "amount2", tv["total_ht"])
            add("total_tva", "amount2", tv["total_tva"])
            add("total_ttc", "amount2", tv["total_ttc"])
            add("net_a_payer", "amount2", tv["net_a_payer"])
        else:
            add("total_credite_ttc", "amount2", tv["total_credite_ttc"])
            for r in tv["refs_facture_origine"]:
                add("refs_facture_origine", "text", r)
        for r in tv["refs_mrn"]:
            add("refs_mrn", "text", r)
        for r in tv["refs_transport"]:
            add("refs_transport", "text", r)
        for i, ln in enumerate(tv["lignes"]):
            add(f"lignes[{i}].libelle", "text", ln["libelle"])
            add(f"lignes[{i}].montant_ht", "amount2", ln["montant_ht"])
            if "montant_ttc" in ln:
                add(f"lignes[{i}].montant_ttc", "amount2", ln["montant_ttc"])
            if k == "ft":
                add(f"lignes[{i}].quantite", "qty", ln["quantite"])
                add(f"lignes[{i}].prix_unitaire", "pu", ln["prix_unitaire"])
                add(f"lignes[{i}].taux_tva", "rate", ln["taux_tva"] if ln["taux_tva"] not in ("0",) else None)
                add(f"lignes[{i}].montant_tva", "amount2", ln["montant_tva"])
                add(f"lignes[{i}].mrn", "text", ln["mrn"])
    return out


def selfcheck(doc, tv, native_bytes, fmt) -> list:
    text = norm_text(extract_text(native_bytes, fmt))
    text_nodot = text.replace(".", "")
    missing = []
    for f, kind, v in checks_for(doc, tv):
        cands = {norm_text(c) for c in _cands(kind, v)}
        if not any(c and (c in text or (kind == "text" and c.replace(".", "") in text_nodot)) for c in cands):
            missing.append(f"{f}={v}")
    return missing


def validate_truth(truth: dict) -> list:
    import jsonschema
    schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
    v = jsonschema.Draft202012Validator(schema)
    return [f"{'/'.join(str(x) for x in e.path)}: {e.message}" for e in v.iter_errors(truth)]
